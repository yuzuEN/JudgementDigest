# -*- coding: utf-8 -*-
"""
structure_tasks.py 及其合併、標準化後處理的離線測試。

之前完全沒有接入 pytest（只有手動執行 `python structure_tasks.py` 時
才會跑 self_check()），PR #4 review 兩次指出這個缺口。這裡先把
self_check() 接進 pytest，再補今天實際修過的兩個 bug 的回歸測試。
"""

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from merge import merge_workbooks  # noqa: E402
from normalize import normalize_law_levels  # noqa: E402
from structure_tasks import (  # noqa: E402
    chinese_number,
    extract_charges,
    extract_laws,
    extract_sentences,
    main_section,
    self_check,
)


class TestSelfCheck(unittest.TestCase):
    def test_self_check_passes(self):
        """把既有的 self_check() 斷言接進 CI，而不是只能手動執行才會跑到。"""
        self_check()


class TestChineseNumberCommaAmount(unittest.TestCase):
    """
    REGRESSION：處罰金金額含千分位逗號（阿拉伯數字寫法）時，SINGLE_RE／
    TOTAL_RE 原本完全比對不到，整案被誤判成「無刑期」（容嫣於 PR #4
    留言回報：extract_sentences('甲犯竊盜罪，處罰金新臺幣30,000元。')
    原本回傳 ([], '無刑期')）。
    """

    def test_comma_amount_is_recognized_as_a_sentence(self):
        declared, executions = extract_sentences("甲犯竊盜罪，處罰金新臺幣30,000元。")
        self.assertEqual(declared, ["罰金新臺幣30,000元"])
        self.assertEqual(executions, [])


class TestLawRESingleCharName(unittest.TestCase):
    """
    REGRESSION：LAW_RE 原本要求法規名稱前綴至少 2 字，「刑法」「民法」這類
    單字全名（不含「法」本身只有 1 字前綴）永遠比對不到（容嫣於 PR #4 留言
    回報：extract_laws('', '刑法第339條第1項；洗錢防制法第19條第1項')
    只回傳洗錢防制法，appendix_parser 的 _LAW_PREFIX 先前已用同樣理由修過）。
    """

    def test_two_character_law_name_is_matched(self):
        laws = extract_laws("", "刑法第339條第1項；洗錢防制法第19條第1項")
        self.assertIn("刑法第339條第1項", laws)
        self.assertIn("洗錢防制法第19條第1項", laws)


class TestChineseNumberMixedDigits(unittest.TestCase):
    """阿拉伯數字混中文單位（今天稍早修的另一個問題）不能被這次改壞。"""

    def test_arabic_digit_run_mixed_with_chinese_unit(self):
        self.assertEqual(chinese_number("10萬"), 100000)
        self.assertEqual(chinese_number("3萬6千"), 36000)


class TestVerdictBoundaries(unittest.TestCase):
    """主文內文字不可被當成後續段落標題，罪數也不屬於罪名。"""

    def test_criminal_facts_column_phrase_does_not_truncate_verdict(self):
        """「犯罪事實欄」是主文內文字，不是新段落。"""
        verdict = "主文：蕭某就起訴書犯罪事實欄一犯詐欺得利罪，處罰金新臺幣參仟元。"
        self.assertIn("詐欺得利罪", main_section(verdict))
        self.assertEqual(extract_charges(verdict, ""), ["詐欺得利罪"])

    def test_criminal_facts_summary_phrase_does_not_truncate_verdict(self):
        """「犯罪事實要旨」同樣不是新段落。"""
        verdict = "法官起立朗讀判決主文、犯罪事實要旨。主文：甲犯竊盜罪，處有期徒刑參月。"
        self.assertEqual(extract_charges(verdict, ""), ["竊盜罪"])

    def test_charge_count_is_not_part_of_charge_name(self):
        """「共貳罪」只是罪數，不得併入罪名。"""
        verdict = "犯詐欺取財罪，共貳罪，各處有期徒刑參月。"
        self.assertEqual(extract_charges(verdict, ""), ["詐欺取財罪"])


class TestLawFallback(unittest.TestCase):
    """理由只含程序法時，必須繼續向事實與適用法條補抓。"""

    def test_procedural_law_does_not_block_substantive_fallback(self):
        """過濾後為空才搜尋下一個來源。"""
        laws = extract_laws(
            "核其所為，係犯刑事訴訟法第299條。",
            "",
            "核其所為，係犯刑法第320條第1項之竊盜罪。",
        )
        self.assertEqual(laws, ["刑法第320條第1項"])


class TestNormalizeLawName(unittest.TestCase):
    """常見法只在清理前綴後從開頭比對，避免子字串誤命中。"""

    def test_enforcement_act_is_not_criminal_code(self):
        """刑法施行法不可正規化成刑法。"""
        self.assertEqual(normalize_law_levels("[中華民國刑法施行法第1條之1]"), ("", ""))

    def test_known_prefixes_are_removed_before_matching(self):
        """正式法名前的版本詞不影響標準化。"""
        self.assertEqual(
            normalize_law_levels("[修正前刑法第320條,前開商標法第95條]"),
            ("[刑法第320條,商標法第95條]", "[刑法第320條,商標法第95條]"),
        )


class TestMergeDeduplication(unittest.TestCase):
    """合併重疊匯出檔時，同工作表的裁判書連結只保留第一筆。"""

    def test_duplicate_judgment_link_is_skipped(self):
        """使用暫存 Excel 驗證去重，不觸碰實際資料檔。"""
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / "a.xlsx", Path(directory) / "b.xlsx"]
            for path, rows in zip(paths, ((["u1", "A"], ["u2", "B"]), (["u2", "B2"], ["u3", "C"]))):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "裁判書資料"
                sheet.append(["裁判書連結", "案件"])
                for row in rows:
                    sheet.append(row)
                workbook.save(path)
            output = Path(directory) / "merged.xlsx"
            self.assertEqual(merge_workbooks(paths, output)["裁判書資料"], 3)
            workbook = load_workbook(output, read_only=True)
            values = list(workbook["裁判書資料"].values)
            self.assertEqual(values[1:], [("u1", "A"), ("u2", "B"), ("u3", "C")])
            workbook.close()


if __name__ == "__main__":
    unittest.main()
