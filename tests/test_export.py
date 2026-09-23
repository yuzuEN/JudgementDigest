# -*- coding: utf-8 -*-
"""
export_excel.py 的離線測試：篩選條件（關鍵字／法院／裁判日期）與 Excel 產出。
全部在暫存 DB 與暫存目錄進行，不碰真正的 judgments.db。
"""

import json
import os
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import export_excel  # noqa: E402
import html_parser  # noqa: E402
from export_excel import BASE_COLUMNS, export_to_excel, fetch_judgments  # noqa: E402

ROWS = [
    # (case_number, court, judgment_date, judgment_type, keyword, verdict)
    ("臺灣臺北地方法院 113 年度訴字第 6910 號民事判決", "臺灣臺北地方法院",
     "2025-01-07", "判決", "臺灣臺北地方法院-民事-判決", "被告應給付原告…"),
    ("臺灣臺北地方法院 113 年度除字第 1836 號民事判決", "臺灣臺北地方法院",
     "2025-01-08", "判決", "臺灣臺北地方法院-民事-判決", "宣告…證券無效"),
    ("臺灣高雄地方法院 112 年度訴字第 1 號民事判決", "臺灣高雄地方法院",
     "2023-06-01", "判決", "借名登記", "原告之訴駁回"),
]


class _TempDbCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        db = os.path.join(self.tmp.name, "t.db")
        self._orig = (export_excel.DB_PATH, html_parser.DB_PATH)
        export_excel.DB_PATH = html_parser.DB_PATH = db
        html_parser.init_db()
        conn = sqlite3.connect(db)
        conn.executemany(
            """INSERT INTO judgments
               (crawl_id, case_number, court, judgment_date, judgment_type,
                keyword, verdict)
               VALUES (?,?,?,?,?,?,?)""",
            [(i, *r) for i, r in enumerate(ROWS, start=1)],
        )
        conn.commit()
        conn.close()

    def tearDown(self):
        export_excel.DB_PATH, html_parser.DB_PATH = self._orig
        self.tmp.cleanup()


class TestFetchJudgments(_TempDbCase):
    def test_no_filter_returns_all(self):
        self.assertEqual(len(fetch_judgments(limit=0)), 3)

    def test_limit(self):
        self.assertEqual(len(fetch_judgments(limit=2)), 2)

    def test_keyword_matches_label(self):
        rows = fetch_judgments(limit=0, keyword="臺灣臺北地方法院-民事-判決")
        self.assertEqual(len(rows), 2)

    def test_keyword_also_matches_case_number_and_verdict(self):
        # -k 同時比對 keyword／裁判字號／主文，維持既有行為
        self.assertEqual(len(fetch_judgments(limit=0, keyword="除字")), 1)
        self.assertEqual(len(fetch_judgments(limit=0, keyword="駁回")), 1)

    def test_court_filter_is_partial_match(self):
        self.assertEqual(len(fetch_judgments(limit=0, court="臺灣臺北")), 2)
        self.assertEqual(len(fetch_judgments(limit=0, court="臺灣高雄地方法院")), 1)

    def test_date_filters_accept_roc_and_ad(self):
        # judgment_date 以西元 ISO 儲存，輸入民國格式也會先正規化
        self.assertEqual(len(fetch_judgments(limit=0, start_date="2025/01/01")), 2)
        self.assertEqual(len(fetch_judgments(limit=0, start_date="114/01/08")), 1)
        self.assertEqual(len(fetch_judgments(limit=0, end_date="2024/12/31")), 1)

    def test_combined_filters(self):
        rows = fetch_judgments(limit=0, court="臺灣臺北",
                               start_date="2025/01/08", end_date="2025/01/08")
        self.assertEqual(len(rows), 1)
        self.assertIn("除字", rows[0]["case_number"])


class TestExportToExcel(_TempDbCase):
    def test_creates_file_with_expected_headers(self):
        from openpyxl import load_workbook
        out = os.path.join(self.tmp.name, "out.xlsx")
        rows = fetch_judgments(limit=0)
        self.assertTrue(export_to_excel(rows, out, include_full_text=False))
        self.assertTrue(os.path.exists(out))

        ws = load_workbook(out).active
        headers = [c.value for c in ws[1]]
        self.assertIn("裁判字號", headers)
        self.assertIn("裁判種類", headers)
        self.assertIn("搜尋關鍵字", headers)
        self.assertNotIn("全文", headers)          # 未加 --full-text
        self.assertEqual(ws.max_row, len(rows) + 1)

    def test_full_text_column_added(self):
        from openpyxl import load_workbook
        out = os.path.join(self.tmp.name, "full.xlsx")
        export_to_excel(fetch_judgments(limit=0), out, include_full_text=True)
        headers = [c.value for c in load_workbook(out).active[1]]
        self.assertIn("全文", headers)

    def test_base_columns_cover_db_schema(self):
        # BASE_COLUMNS 的每個欄位都必須存在於 judgments 表，否則匯出會整批失敗
        conn = sqlite3.connect(export_excel.DB_PATH)
        cols = {r[1] for r in conn.execute("PRAGMA table_info(judgments)")}
        conn.close()
        for db_col, _label in BASE_COLUMNS:
            self.assertIn(db_col, cols, msg=db_col)


class ApplicableLawsTextTests(unittest.TestCase):
    """
    「適用法條」易讀欄是 applicable_laws_json 的呈現，不是另一套抽取規則。
    這些測試鎖住兩件事：渲染忠於 JSON，以及壞輸入不會讓整批匯出失敗。
    """

    def test_joins_keys_in_order(self):
        row = {"applicable_laws_json": json.dumps([
            {"key": "民法§767第1項前段"}, {"key": "民法§470第2項"},
        ], ensure_ascii=False)}
        self.assertEqual(export_excel._applicable_laws_text(row),
                         "民法§767第1項前段；民法§470第2項")

    def test_deduplicates_preserving_order(self):
        # 同一條文常在主文與理由各出現一次，重複輸出只會讓欄位更難讀
        row = {"applicable_laws_json": json.dumps([
            {"key": "民法§179"}, {"key": "民法§767"}, {"key": "民法§179"},
        ], ensure_ascii=False)}
        self.assertEqual(export_excel._applicable_laws_text(row), "民法§179；民法§767")

    def test_keeps_sub_article_and_paragraph(self):
        # 這正是原始 applicable_laws 欄會截斷成「第436條之1第3」的情形
        row = {"applicable_laws_json": json.dumps(
            [{"key": "民事訴訟法§436之1第3項"}], ensure_ascii=False)}
        self.assertEqual(export_excel._applicable_laws_text(row), "民事訴訟法§436之1第3項")

    def test_falls_back_to_legacy_column(self):
        """REGRESSION：舊資料庫尚未回填時，適用法條欄不可整欄空白。

        BASE_COLUMNS 改成只讀 applicable_laws_json 之後，還沒跑過
        backfill_structured.py 的資料庫（例如刑事那邊現有的一萬多筆）
        匯出後這一欄會全空，而且沒有任何警告——合併前它是有資料的。
        """
        row = {"applicable_laws": "民法第179條、第184條", "applicable_laws_json": ""}
        self.assertEqual(export_excel._applicable_laws_text(row),
                         "民法第179條、第184條（未結構化）")
        # JSON 壞掉時同樣退回原始欄
        self.assertEqual(
            export_excel._applicable_laws_text(
                {"applicable_laws": "民法第5條", "applicable_laws_json": "{壞掉"}),
            "民法第5條（未結構化）")
        # 兩者皆空才是空字串
        self.assertEqual(
            export_excel._applicable_laws_text(
                {"applicable_laws": "", "applicable_laws_json": ""}), "")

    def test_bad_input_returns_empty_not_raises(self):
        for bad in (None, "", "[]", "not json", "{}", json.dumps([1, 2])):
            with self.subTest(bad=bad):
                self.assertEqual(
                    export_excel._applicable_laws_text({"applicable_laws_json": bad}), "")

    def test_law_columns_are_adjacent(self):
        # 易讀欄與 JSON 欄若相隔數十欄，讀表的人得左右橫跨整張表才能對照
        labels = [l for _, l in export_excel.STRUCTURED_EXPORT_COLUMNS]
        self.assertEqual(
            labels[labels.index("主要法規"):labels.index("主要法規") + 4],
            ["主要法規", "法條引用數", "適用法條", "適用法條（結構化JSON）"])

    def test_raw_applicable_laws_not_exported(self):
        # 原始欄 85.7% 被截斷且無獨佔資訊，留在 Excel 裡只會誤導
        self.assertNotIn("applicable_laws", [c for c, _ in BASE_COLUMNS])


if __name__ == "__main__":
    unittest.main()
