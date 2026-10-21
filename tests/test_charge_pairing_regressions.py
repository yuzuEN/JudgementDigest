"""PR #12：主罪配對、連續項款，以及不明酒／毒的不能安全駕駛。"""
import json
import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook

from build_structured_review import HEADERS, build_row
from normalize import normalize_charge, normalize_charges, normalize_workbook
from structure_tasks import matched_laws, resolve_law_references


def review(charge, reason='', facts='', *, use_fallback=True):
    row, _ = build_row({'被告／相對人': '甲某', '主文': f'甲某犯{charge}，處有期徒刑參月。',
                        '理由': reason, '事實': facts}, use_fallback=use_fallback)
    return dict(zip(HEADERS, row))


class TestChargeFirstPairing(unittest.TestCase):
    def test_real_court_conclusion_variants(self):
        # judgments.id=68780（107簡88）、69004（106審易2718）的法院原句。
        for reason in (
            '核被告所為，係違反毒品危害防制條例第10條第2項施用第二級毒品罪。',
            '三、論罪科刑（一）核被告所為，是犯毒品危害防制條例第10條第2項之施用第二級毒品罪。',
        ):
            with self.subTest(reason=reason):
                row = review('施用第二級毒品罪', reason)
                self.assertEqual(json.loads(row['法條（對應原始）']), [['毒品危害防制條例第10條第2項']])

    def test_real_court_and_prosecution_in_same_sentence(self):
        # 77748（107簡2258）與73711（107簡1548）；後者姓名以甲某代替。
        cases = [
            ('詐欺取財罪', '核被告所為，係犯刑法第339條第1項之詐欺取財罪，公訴意旨認被告係同時觸犯詐欺取財罪及詐欺得利罪，容有未當。',
             ['刑法第339條第1項']),
            ('強制罪', '二、核被告甲某所為，係犯刑法第304條之強制罪，本件公訴意旨原認被告所犯強制與傷害二罪間有想像競合關係，因告訴人就傷害罪部分撤回刑事告訴，故僅就強制罪部分予以論罪科刑，附此敘明。',
             ['刑法第304條']),
        ]
        for charge, reason, expected in cases:
            with self.subTest(charge=charge):
                self.assertEqual(json.loads(review(charge, reason)['法條（對應原始）']), [expected])

    def test_real_traffic_long_names_and_article_part(self):
        # 68624（107審交簡26）精確配對；68930（107審交簡4）唯一粗分類配對。
        cases = [
            ('汽車駕駛人，無駕駛執照駕車，因過失傷害致人受傷罪',
             '係犯道路交通管理處罰條例第86條第1項、刑法第284條第1項前段之汽車駕駛人，無駕駛執照駕車，因過失傷害致人受傷罪，且應依道路交通管理處罰條例第86條第1項之規定，加重其刑。'),
            ('過失傷害罪',
             '是核被告所為，係犯道路交通管理處罰條例第86條第1項、刑法第284條第1項前段之汽車駕駛人，無駕駛執照駕車，行經行人穿越道不依規定讓行人優先通行，因過失傷害致人傷害罪，並應依道路交通管理處罰條例第86條第1項規定加重其刑。'),
        ]
        for charge, reason in cases:
            with self.subTest(charge=charge):
                self.assertEqual(json.loads(review(charge, reason)['法條（對應原始）']),
                                 [['道路交通管理處罰條例第86條第1項', '刑法第284條第1項']])

    def test_rejected_prosecution_after_court_does_not_supply_another_charge(self):
        reason = ('核被告所為，係犯刑法第339條第1項之詐欺取財罪，'
                  '公訴意旨認被告係犯刑法第339條之3第1項之電腦詐欺罪，尚有未洽。')
        self.assertEqual(json.loads(review('詐欺取財罪', reason)['法條（對應原始）']), [['刑法第339條第1項']])
        self.assertEqual(json.loads(review('電腦詐欺罪', reason)['法條（對應原始）']), [[]])

    def test_explicit_court_resumption_after_prosecution(self):
        reason = ('公訴意旨認被告係犯刑法第339條之3之電腦詐欺罪，'
                  '惟本院認核被告所為，係犯刑法第339條第1項之詐欺取財罪。')
        self.assertEqual(json.loads(review('詐欺取財罪', reason)['法條（對應原始）']), [['刑法第339條第1項']])

    def test_prosecution_and_acquittal_never_become_supporting_charges(self):
        for reason in ('公訴意旨認被告係違反毒品危害防制條例第10條第2項施用第二級毒品罪。',
                       '核被告所為，並不成立刑法第320條之竊盜罪，應諭知無罪。'):
            self.assertEqual(json.loads(review('竊盜罪', reason)['法條（對應原始）']), [[]])
            from structure_tasks import extract_supporting_charges
            self.assertEqual(extract_supporting_charges(reason), [])

    def test_article_part_and_long_name_stop_before_other_charge(self):
        reason = ('係犯刑法第284條第1項前段之汽車駕駛人，因過失傷害致人受傷罪、'
                  '刑法第185條之4之肇事逃逸罪。')
        self.assertEqual(json.loads(review('過失傷害罪', reason)['法條（對應原始）']), [['刑法第284條第1項']])
        self.assertEqual(json.loads(review('肇事逃逸罪', reason)['法條（對應原始）']), [['刑法第185條之4']])

    def test_exact_match_in_later_source_precedes_coarse_match(self):
        laws, source = matched_laws('過失傷害罪', '過失傷害罪', '',
                                   '係犯刑法第284條第1項之汽車駕駛人，因過失傷害致人受傷罪。',
                                   '係犯刑法第284條之過失傷害罪。', '', normalize_charge)
        self.assertEqual((laws, source), (['刑法第284條'], '事實'))

    def test_real_common_offender_wording(self):
        # 136705（111審訴2210）。
        reason = ('是核被告所為，係犯刑法第339條之4第1項第2款之三人以上共同犯詐欺取財罪及'
                  '洗錢防制法第2條第2款隱匿特定犯罪所得之去向而依同法第14條第1項論處之洗錢罪。'
                  '公訴意旨認被告係犯刑法第339條第1項之詐欺取財罪嫌，依上開說明，容有未恰。')
        self.assertEqual(json.loads(review('三人以上共同詐欺取財罪', reason)['法條（對應原始）']),
                         [['刑法第339條之4第1項第2款']])

    def test_suspect_suffix_requires_exact_main_charge_not_coarse_guess(self):
        reason = '二、核被告甲某所為，係犯刑法第354條之毀損罪嫌。'
        self.assertEqual(json.loads(review('毀損罪', reason)['法條（對應原始）']), [['刑法第354條']])
        # 精確配對的邊界不變；正式備援的較弱來源另由 integration tests 記錄。
        self.assertEqual(json.loads(review('毀損他人物品罪', reason, use_fallback=False)['法條（對應原始）']), [[]])

    def test_prior_history_in_preface_does_not_delete_current_court_finding(self):
        # 60619（107訴636）：前案是本案持有期間的起點，非前案定罪引用。
        reason = ('是核被告自前案遭查獲後起至本案遭查獲時止，持有本案槍枝之行為，'
                  '係犯槍砲彈藥刀械管制條例第8條第4項之非法持有可發射子彈具殺傷力之槍枝罪。')
        self.assertEqual(json.loads(review('非法持有可發射子彈具殺傷力之槍枝罪', reason)['法條（對應原始）']),
                         [['槍砲彈藥刀械管制條例第8條第4項']])

    def test_death_wording_is_equivalent_but_attempt_is_not(self):
        row = review('業務過失致人於死罪', '核被告所為，係犯刑法第276條第2項之業務過失致死罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [['刑法第276條第2項']])
        row = review('逾越牆垣竊盜罪', '核被告所為，係犯刑法第321條第2項、第1項第2款之逾越牆垣竊盜未遂罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [[]])

    def test_same_law_set_under_different_names_is_not_law_ambiguity(self):
        # 69784（107審交簡39）：兩種法院罪名寫法引用相同的條項。
        reason = ('核其所為，係犯道路交通管理處罰條例第86條第1項、刑法第284條第1項前段之'
                  '汽車駕駛人，酒醉駕車，因過失傷害致人受傷罪。'
                  '核其所為，係犯道路交通管理處罰條例第86條第1項、刑法第284條第1項前段之'
                  '汽車駕駛人酒醉駕車，而犯過失傷害罪。'
                  '起訴書意旨認被告係犯刑法第284條第1項前段之過失傷害罪，顯有誤認。')
        self.assertEqual(json.loads(review('過失傷害罪', reason)['法條（對應原始）']),
                         [['道路交通管理處罰條例第86條第1項', '刑法第284條第1項']])

    def test_helping_offender_wording_keeps_property_and_benefit_separate(self):
        # 84444（113審簡606）：主文「幫助犯詐欺得利罪」不取另一取財罪的第1項。
        reason = ('核被告所為，係犯刑法第30條第1項前段、第339條第1項之幫助犯詐欺取財罪及'
                  '刑法第30條第1項前段、第339條第2項之幫助犯詐欺得利罪。')
        self.assertEqual(json.loads(review('詐欺得利罪', reason)['法條（對應原始）']),
                         [['刑法第30條第1項', '刑法第339條第2項']])

    def test_rejected_summary_indictment_is_not_a_second_court_candidate(self):
        reason = ('核被告所為，係犯道路交通管理處罰條例第86條第1項、刑法第284條第1項之'
                  '汽車駕駛人，未依規定讓行人優先通行，過失傷害罪。'
                  '聲請簡易判決處刑書雖認被告僅係犯刑法第284條第1項前段之過失傷害罪，法條容有未洽。')
        self.assertEqual(json.loads(review('過失傷害人罪', reason)['法條（對應原始）']),
                         [['道路交通管理處罰條例第86條第1項', '刑法第284條第1項']])

    def test_supporting_charge_count_cannot_replace_a_different_main_offense(self):
        # 82023 的失敗模式：只有罪數相同不足以把主文施用改成理由的販賣。
        values, _ = build_row({'被告／相對人': '甲某',
            '主文': '甲某施用第一級毒品，處有期徒刑柒月。又施用第二級毒品，處有期徒刑肆月。',
            '理由': '核被告所為，係犯毒品危害防制條例第4條第1項之販賣第一級毒品罪。'
                    '另係犯毒品危害防制條例第4條第2項之販賣第二級毒品罪。'})
        row = dict(zip(HEADERS, values))
        self.assertEqual(json.loads(row['罪名（原始）']), ['施用第一級毒品', '施用第二級毒品'])
        self.assertEqual(json.loads(row['法條（對應原始）']), [[], []])

    def test_generic_public_danger_requires_driving_article_and_consistent_subclause(self):
        main = '駕駛動力交通工具，而有吐氣所含酒精濃度達每公升零點二五毫克以上情形'
        for article, expected in [('185條之3第1項第1款', [['刑法第185條之3第1項第1款']]),
                                  ('173條第1項', [[]]), ('185條之3第1項第3款', [[]])]:
            values, _ = build_row({'被告／相對人': '甲某', '主文': f'甲某{main}，處有期徒刑參月。',
                                  '理由': f'核被告所為，係犯刑法第{article}之公共危險罪。'})
            row = dict(zip(HEADERS, values))
            self.assertEqual(json.loads(row['罪名（原始）']), [main])
            self.assertEqual(json.loads(row['法條（對應原始）']), expected)

    def test_helping_offense_is_conviction_and_law_name_is_clean(self):
        row = review('詐欺取財罪', '核被告所為，係幫助犯刑法第339條第1項之詐欺取財罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [['刑法第339條第1項']])

    def test_continuous_paragraphs_and_attempt(self):
        row = review('幫助販賣第三級毒品未遂罪',
                     '核被告所為，係犯刑法第30條、毒品危害防制條例第4條第3項、第6項之幫助販賣第三級毒品未遂罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [
            ['刑法第30條', '毒品危害防制條例第4條第3項', '毒品危害防制條例第4條第6項']])

    def test_continuous_items_do_not_recover_other_offense(self):
        row = review('三人以上共同冒用政府機關名義詐欺取財罪',
                     '核被告所為，係犯刑法第216條、第211條之行使偽造公文書罪、'
                     '刑法第339條之4第1項第2款、第1款之三人以上共同冒用政府機關名義詐欺取財罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [
            ['刑法第339條之4第1項第2款', '刑法第339條之4第1項第1款']])

    def test_parenthesis_after_charge_and_other_victim(self):
        row = review('對直系血親尊親屬重傷害未遂罪',
                     '係犯刑法第306條第1項之侵入住宅罪、刑法第280條、第278條第3項、第1項之'
                     '對直系血親尊親屬重傷害未遂罪（告訴人乙部分）、刑法第278條第3項、第1項之重傷害未遂罪（丙部分）。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [
            ['刑法第280條', '刑法第278條第3項', '刑法第278條第1項']])

    def test_court_charge_precedes_rejected_prosecution_theory(self):
        row = review('詐欺得利罪',
                     '核被告所為，係刑法第339條第1項詐欺得利罪之幫助犯。'
                     '公訴意旨認被告係犯刑法第339條之3第1項之幫助電腦詐欺罪，顯屬速斷。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [['刑法第339條第1項']])

    def test_no_case_fallback_for_another_offense(self):
        laws, _ = matched_laws('竊盜罪', '竊盜罪', '',
                               '係犯刑法第339條第1項之詐欺取財罪。', '', '刑法第339條', normalize_charge)
        self.assertEqual(laws, [])

    def test_same_coarse_class_does_not_merge_two_named_offenses(self):
        row = review('竊盜罪', '係犯刑法第320條第1項之普通竊盜罪、刑法第321條第1項之攜帶兇器竊盜罪。')
        self.assertEqual(json.loads(row['法條（對應原始）']), [[]])

    def test_item_expansion_stops_at_description_and_sentence(self):
        text = '刑法第339條第1項之詐欺罪。第2項。'
        self.assertEqual(resolve_law_references(text, {'刑法'}), text)


class TestDrivingEvidence(unittest.TestCase):
    def assert_class(self, expected, reason='', facts='', charge='不能安全駕駛動力交通工具罪'):
        row = review(charge, reason, facts)
        self.assertEqual(json.loads(row['罪名（標準化）']), [expected])
        self.assertEqual(json.loads(row['罪名（原始）']), [charge])

    def test_generic_name_is_not_automatically_alcohol(self):
        self.assertEqual(normalize_charge('不能安全駕駛動力交通工具罪'), '')
        self.assert_class('')

    def test_explicit_substance_names(self):
        for charge, expected in [('施用安非他命致不能安全駕駛動力交通工具罪', '毒駕罪'),
                                 ('服用酒類致不能安全駕駛動力交通工具罪', '酒駕罪')]:
            self.assertEqual(normalize_charge(charge), expected)

    def test_explicit_article_subclause_in_main_charge(self):
        for numeral in ['185', '一百八十五', '一八五']:
            self.assertEqual(normalize_charge(f'刑法第{numeral}條之3第1項第3款之不能安全駕駛動力交通工具罪'), '毒駕罪')

    def test_reason_bound_subclauses(self):
        for item, expected in [(1, '酒駕罪'), (2, '酒駕罪'), (3, '毒駕罪'), (4, '毒駕罪')]:
            self.assert_class(expected, f'核被告所為，係犯刑法第185條之3第1項第{item}款之不能安全駕駛動力交通工具罪。')

    def test_causal_drug_facts(self):
        self.assert_class('毒駕罪', facts='甲施用甲基安非他命後，仍駕駛汽車上路。')

    def test_causal_alcohol_facts(self):
        self.assert_class('酒駕罪', facts='甲飲用啤酒後，仍騎乘機車上路。')

    def test_actual_alcohol_reading(self):
        self.assert_class('酒駕罪', facts='甲騎乘機車遭查獲，測得吐氣所含酒精濃度達每公升1.45毫克。')

    def test_other_drug_offense_is_not_drug_driving_evidence(self):
        self.assert_class('', '另係犯毒品危害防制條例第10條第2項之施用第二級毒品罪。',
                          '甲施用安非他命。翌日甲駕車遭查獲。')

    def test_prior_case_and_legal_discussion_are_not_current_evidence(self):
        self.assert_class('', facts='甲曾因施用安非他命後駕車而受判決。刑法規定施用毒品後不能安全駕駛應予處罰。')

    def test_zero_alcohol_does_not_conflict_with_drug(self):
        self.assert_class('毒駕罪', facts='甲施用安非他命後駕車，吐氣酒精濃度為0.00毫克。')

    def test_denied_drug_use_is_not_positive_evidence(self):
        self.assert_class('', facts='甲否認施用安非他命後駕車。')

    def test_multisentence_law_quotation_does_not_create_false_drug_conflict(self):
        self.assert_class('酒駕罪', facts='甲飲酒後駕車肇事。相關條文規定：「三、酒醉駕車。'
                          '四、吸食毒品、迷幻藥或麻醉藥品駕車。」')

    def test_facts_alone_do_not_assign_both_without_conviction(self):
        self.assert_class('', facts='甲施用安非他命及飲用啤酒後，仍駕車上路。')

    def assert_mixed(self, reason='', charge='不能安全駕駛動力交通工具罪'):
        row = review(charge, reason)
        self.assertEqual(json.loads(row['罪名（標準化）']), ['酒駕罪', '毒駕罪'])
        self.assertEqual(json.loads(row['罪名（對應標準化）']), [['酒駕罪', '毒駕罪']])
        self.assertEqual(len(json.loads(row['罪名（原始）'])), 1)
        self.assertEqual(len(json.loads(row['宣告主刑'])), 1)
        self.assertEqual(len(json.loads(row['法條（對應原始）'])), 1)
        return row

    def test_mixed_subclauses_preserve_one_offense_and_sentence(self):
        self.assert_mixed('係犯刑法第185條之3第1項第1款、第3款之不能安全駕駛動力交通工具罪。')
        row = self.assert_mixed(charge='刑法第一八五條之三第一項第一款、第三款之不能安全駕駛動力交通工具罪')
        self.assertEqual(json.loads(row['法條（對應原始）']), [[
            '刑法第一八五條之三第一項第一款', '刑法第一八五條之三第一項第三款']])

    def test_mixed_conduct_in_conviction_section(self):
        self.assert_mixed('二、論罪科刑：被告飲用啤酒及施用安非他命後，仍駕車上路。')

    def test_same_words_in_generic_discussion_do_not_imply_mixed(self):
        self.assert_class('', '二、論罪科刑：被告明知飲酒或施用安非他命後不得駕車。')
        self.assert_class('', '二、論罪科刑：刑法規定：「飲酒或施用安非他命後不得駕車。」')
        self.assert_class('酒駕罪', '二、論罪科刑：被告飲酒後駕車。另係犯施用毒品罪。')

    def test_different_conviction_groups_are_not_one_mixed_offense(self):
        self.assert_class('', '係犯刑法第185條之3第1項第1款之不能安全駕駛罪、'
                          '刑法第185條之3第1項第3款之不能安全駕駛罪。')

    def test_contradictory_subclause_and_name_stay_unresolved(self):
        self.assert_class('', charge='刑法第185條之3第1項第1款之服用毒品致不能安全駕駛罪')
        self.assert_class('', '係犯刑法第185條之3第1項第1款之服用毒品致不能安全駕駛罪。')

    def test_mixed_labels_reach_model_export_without_duplicate_sentence(self):
        from refresh_modeldata import apply_review
        row = self.assert_mixed('係犯刑法第185條之3第1項第1款、第3款之不能安全駕駛動力交通工具罪。')
        result = apply_review({'fact': '固定事實', 'meta': {}}, row, 'reextracted', 'exact_url')
        self.assertEqual(result['meta']['accusation'], ['酒駕罪', '毒駕罪'])
        self.assertEqual(result['meta']['reextraction']['standard_charges_by_offense'], [['酒駕罪', '毒駕罪']])
        self.assertTrue(result['meta']['reextraction']['task_validity']['accusation'])
        self.assertEqual(result['meta']['term_of_imprisonment']['imprisonment'], 3)
        row['罪名（原始）'] = '["不能安全駕駛罪","未分類罪"]'
        row['罪名（對應標準化）'] = '[["酒駕罪","毒駕罪"],[]]'
        result = apply_review({'fact': '固定事實', 'meta': {}}, row, 'reextracted', 'exact_url')
        self.assertFalse(result['meta']['reextraction']['task_validity']['accusation'])

    def test_legacy_mixed_labels_are_flat(self):
        self.assertEqual(normalize_charges('[不能安全駕駛罪]',
            '係犯刑法第185條之3第1項第1款、第3款之不能安全駕駛罪。'), '[酒駕罪,毒駕罪]')

    def test_known_alcohol_charge_not_changed_by_separate_drug_offense(self):
        self.assert_class('酒駕罪', '係犯毒品危害防制條例第10條第2項之施用第二級毒品罪。',
                          charge='服用酒類致不能安全駕駛動力交通工具罪')

    def test_legacy_normalization_can_use_context_without_changing_raw_label(self):
        self.assertEqual(normalize_charges('[不能安全駕駛動力交通工具罪]',
            '係犯刑法第185條之3第1項第3款之不能安全駕駛動力交通工具罪。'), '[毒駕罪]')

    def test_multi_charge_context_keeps_existing_all_or_nothing_policy(self):
        charges = '[不能安全駕駛動力交通工具罪,竊盜罪]'
        self.assertEqual(normalize_charges(charges), '')
        self.assertEqual(normalize_charges(charges,
            '核被告所為，係犯刑法第185條之3第1項第1款之不能安全駕駛動力交通工具罪。'),
            '[酒駕罪,竊盜罪]')
        self.assertEqual(normalize_charges('[不能安全駕駛罪,未分類罪]',
            facts='被告飲酒後駕車。'), '')

    def test_real_ambiguous_driving_conclusions_reach_both_interfaces(self):
        # 原句：41252（109交簡553）、83390（113交簡211），未改罪名或款別。
        for reason, expected in [
            ('㈠核被告所為，係犯刑法第185條之3第1項第1款之不能安全駕駛動力交通工具罪', '酒駕罪'),
            ('(二)核被告所為，係犯刑法第185條之3第1項第4款之不能安全駕駛動力交通工具罪', '毒駕罪'),
        ]:
            self.assert_class(expected, reason=reason)
            self.assertEqual(normalize_charges('[不能安全駕駛動力交通工具罪]', reason=reason),
                             f'[{expected}]')

    def test_explicit_names_do_not_read_unrelated_context(self):
        for charge, label in [('酒駕罪', '酒駕罪'), ('毒駕罪', '毒駕罪'),
                              ('酒後不能安全駕駛致人於死罪', '酒駕罪'),
                              ('施用毒品不能安全駕駛致重傷罪', '毒駕罪')]:
            self.assertEqual(normalize_charges(f'[{charge},竊盜罪]',
                '檢察官主張被告施用毒品後駕車。', '被告飲酒後駕車。'), f'[{label},竊盜罪]')

    def test_prosecution_and_rejected_conduct_are_not_factual_evidence(self):
        for text in ['檢察官主張被告施用安非他命後駕車。',
                     '聲請簡易判決處刑意旨認被告飲酒後駕車。',
                     '被告施用安非他命後駕車，尚難認定。',
                     '被告飲酒後駕車之犯罪不成立。']:
            with self.subTest(text=text):
                self.assert_class('', reason=text)
                self.assertEqual(normalize_charges('[不能安全駕駛罪]', facts=text), '')

    def test_prosecution_mixed_conduct_is_not_court_mixed_finding(self):
        self.assert_class('', '二、論罪科刑：檢察官主張被告飲用啤酒及施用安非他命後駕車。')

    def test_context_attachment_is_not_court_evidence(self):
        text = ('附件：臺灣臺北地方檢察署檢察官起訴書110年度被告甲'
                '係犯刑法第185條之3第1項第3款之不能安全駕駛罪。')
        self.assertEqual(normalize_charges('[不能安全駕駛罪]', reason=text), '')

    def test_same_case_separate_drug_conviction_does_not_change_driving(self):
        self.assert_class('酒駕罪',
            '核被告所為，係犯刑法第185條之3第1項第1款之不能安全駕駛動力交通工具罪。'
            '另係犯毒品危害防制條例第10條第2項之施用第二級毒品罪。',
            '被告施用安非他命。')

    def test_mixed_group_cannot_override_contradictory_pair(self):
        self.assert_class('',
            '係犯刑法第185條之3第1項第1款之服用毒品致不能安全駕駛罪。'
            '係犯刑法第185條之3第1項第1款、第3款之不能安全駕駛罪。')

    def test_driving_result_must_match_context_offense(self):
        self.assert_class('', '係犯刑法第185條之3第1項第1款之不能安全駕駛罪。',
                          charge='不能安全駕駛致人於死罪')
        self.assert_class('酒駕罪', '係犯刑法第185條之3第2項之酒後不能安全駕駛致死罪。',
                          charge='不能安全駕駛致人於死罪')
        self.assert_class('毒駕罪', '係犯刑法第185條之3第2項之施用毒品不能安全駕駛致重傷罪。',
                          charge='不能安全駕駛致重傷罪')
        self.assert_class('', '係犯刑法第185條之3第2項之施用毒品不能安全駕駛致重傷罪。',
                          charge='不能安全駕駛致死罪')
        self.assert_class('酒駕罪', '係犯刑法第185條之3第3項之不能安全駕駛動力交通工具致人於死罪。',
                          facts='被告飲酒後駕車肇事。', charge='不能安全駕駛動力交通工具致人於死罪')

    def test_multiple_driving_offenses_need_pairing_not_shared_facts(self):
        self.assertEqual(normalize_charges('[酒駕罪,不能安全駕駛罪]',
            facts='被告飲酒後駕車。'), '')
        self.assertEqual(normalize_charges('[酒駕罪,不能安全駕駛罪]',
            reason='係犯刑法第185條之3第1項第3款之不能安全駕駛罪。'), '[酒駕罪,毒駕罪]')

    def test_real_driving_obstruction_does_not_borrow_alcohol_label(self):
        # 84235（113審交簡118）主文兩罪，姓名以甲某代替；節錄論罪配對。
        from refresh_modeldata import apply_review
        values, _ = build_row({
            '被告／相對人': '甲某',
            '主文': '甲某吐氣所含酒精濃度達每公升零點二五毫克以上而駕駛動力交通工具，'
                    '處有期徒刑伍月。又犯駕駛動力交通工具妨害公務執行罪，處有期徒刑陸月。',
            '理由': '係犯刑法第185條之3第1項第1款之吐氣所含酒精濃度達0.25毫克以上'
                    '而駕駛動力交通工具罪。',
            '事實': '被告飲酒後騎車為警查獲。',
        })
        row = dict(zip(HEADERS, values))
        self.assertEqual(json.loads(row['罪名（對應標準化）']), [['酒駕罪'], []])
        exported = apply_review({'fact': '固定事實', 'meta': {}}, row, 'reextracted', 'id')
        self.assertEqual(exported['meta']['accusation'], [])
        self.assertEqual(exported['meta']['reextraction']['partial_accusation'], ['酒駕罪'])
        self.assertFalse(exported['meta']['reextraction']['task_validity']['accusation'])

    def test_workbook_separates_attachment_field_before_merging(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.xlsx', Path(directory) / 'output.xlsx'
            book = Workbook()
            raw = book.active
            raw.title = '裁判書資料'
            raw.append(['裁判字號', '裁判書連結', '事實及理由', '理由'])
            raw.append(['甲案', 'url-a',
                '附件：臺灣臺北地方檢察署檢察官起訴書110年度被告甲'
                '係犯刑法第185條之3第1項第3款之不能安全駕駛罪。',
                '核被告所為，係犯刑法第185條之3第1項第1款之不能安全駕駛罪。'])
            sheet = book.create_sheet('罪法刑結構化')
            sheet.append(['裁判字號', '裁判書連結', '法條', '罪名', '所有宣告刑', '應執行刑'])
            sheet.append(['甲案', 'url-a', '', '[不能安全駕駛罪,竊盜罪]', '[]', '[]'])
            book.save(source)
            book.close()
            normalize_workbook(source, output)
            result = load_workbook(output)
            try:
                self.assertEqual(result['標準化']['F2'].value, '[酒駕罪,竊盜罪]')
            finally:
                result.close()

    def test_workbook_uses_only_uniquely_matched_source_row(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.xlsx'
            output = Path(directory) / 'output.xlsx'
            book = Workbook()
            raw = book.active
            raw.title = '裁判書資料'
            raw.append(['裁判字號', '裁判書連結', '理由'])
            reason = '係犯刑法第185條之3第1項第3款之不能安全駕駛動力交通工具罪。'
            raw.append(['甲案', 'url-a', reason])
            raw.append(['乙案', 'url-b', reason])
            raw.append(['乙案', 'url-b', reason])
            sheet = book.create_sheet('罪法刑結構化')
            sheet.append(['裁判字號', '裁判書連結', '法條', '罪名', '所有宣告刑', '應執行刑'])
            for case, url in [('甲案', 'url-a'), ('乙案', 'url-b'), ('丙案', 'url-c')]:
                sheet.append([case, url, '[刑法第185條之3]', '[不能安全駕駛動力交通工具罪]', '[]', '[]'])
            book.save(source)
            book.close()
            normalize_workbook(source, output)
            result = load_workbook(output)
            try:
                self.assertEqual(result['標準化']['F2'].value, '[毒駕罪]')
                self.assertIsNone(result['標準化']['F3'].value)
                self.assertIsNone(result['標準化']['F4'].value)
            finally:
                result.close()


if __name__ == '__main__':
    unittest.main()
