"""從裁判書 Excel 萃取每案的定罪法條、罪名、宣告刑與應執行刑。"""

import argparse
import re
from pathlib import Path
from typing import Optional, Union

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from structuring import LAW_ALIASES, LAW_WHITELIST
from driving_charge import is_driving_charge, driving_charge_kind, driving_charge_kinds, driving_fact_kinds, strip_law_quotations, mixed_driving_in_conviction, explicitly_mixed_driving


SOURCE_SHEET = "裁判書資料"
OUTPUT_SHEET = "罪法刑結構化"
NON_GUILTY_WORDS = ("無罪", "不受理", "免訴")
NUM = "零〇○一二兩三四五六七八九十百千萬億壹貳參叁叄肆伍陸柒捌玖拾佰仟0123456789"
# {1,18} 允許單字法律名（民法、刑法）：{2,18} 會要求「法／條例」前至少 2 字，
# 「民法」「刑法」只有 1 字前綴，永遠比對不到（容嫣於 PR #4 留言回報，appendix_parser
# 的 _LAW_PREFIX 已用同樣理由改過）。lazy 量詞會優先取最短，不會因此多抓到
# 「應依」等觸發詞殘留。
LAW_RE = re.compile(
    rf"(?:係犯|共同犯|違反|適用|應依|依|按|犯|準用|及|與|暨|、|，|；|^)"
    rf"([^，、；。：（）()\s]{{1,18}}?(?:條例|法))第[{NUM}]+條"
    rf"(?:之[{NUM}]+)?(?:第[{NUM}]+項)?(?:第[{NUM}]+款)?"
)
LAW_ID_RE = re.compile(rf"^(.+?(?:條例|法))第([{NUM}]+)條(?:之([{NUM}]+))?")
CHARGE_AFTER_LAW_RE = re.compile(r"之?([^，、；。（）()]{2,80}?罪)(?=之幫助犯|[，、；。及與暨（）()]|$)")
VERDICT_CHARGE_RE = re.compile(
    r"(?<!罪)犯(?!罪事實)([^；。]{2,160}?罪)"
    rf"(?=，(?:共?[{NUM}]+罪，)?(?:(?:(?:均)?累犯|未遂)，)?(?:均處|各處|處|科|免刑)|[；。]|$)"
)
CHARGE_COUNT_SUFFIX_RE = re.compile(rf"，(?:(?:均)?累犯，)?共?[{NUM}]+罪$")
APPENDIX_REFERENCE_RE = re.compile(r"^(?:犯)?(?:(?:如|本).{0,12})?附表")
MAIN_END_RE = re.compile(r"(?:犯罪事實及理由|事實及理由|事實與理由|犯罪事實)(?!欄|要旨)")
CONCLUSION_CHARGE_RE = re.compile(r"(?:分別係犯|分別犯|[係是](?:幫助|共同)?犯|係違反|係(?=刑法第)|應論被告|各犯(?!行(?:時|為|[，。；、]|$))|均犯)")
PROSECUTION_CLAIM_RE = re.compile(r"公訴意旨|起訴(?:書)?意旨|聲請(?:簡易判決處刑)?意旨|聲請簡易判決處刑書(?:雖)?認|檢察官(?:認為|認|主張)")
COURT_RESUMPTION_RE = re.compile(r"(?:惟|然|但)[，,]?(?:本院|核被告|核其所為)|本院(?:認為|認|查)")
REJECTED_CONVICTION_RE = re.compile(r"不另為無罪|不構成|不成立|尚難認|無罪")
VERSION_COMPARISON_CLAUSE_RE = re.compile(
    r'(?:於)?(?:修正|修法)[前後](?:被告)?(?:所犯為|係犯|係適用)[^，。；]*'
)
SENTENCING_CITATION_TAIL_RE = re.compile(
    r'^(?:前段|後段)?(?:之|的)?規定[，,]?'
    r'(?:就被告犯行)?(?:應予|予以|酌量|酌予)?'
    r'(?:加重|減輕(?:或免除)?|酌減|免除)(?:其刑|被告(?:之|的)刑度|其刑度)'
)
LAW_POINT_RE = re.compile(
    r"(?=(?:[㈠㈡㈢㈣㈤㈥㈦㈧㈨㈩]|[（(][一二三四五六七八九十]+[）)]|"
    r"[一二三四五六七八九十]+、))"
)
LAW_REFERENCE_NAMES = (
    "同條例", "前開條例", "上述條例", "上開條例", "該條例", "本條例", "前揭條例", "前述條例",
    "同法", "前開法", "上述法", "上開法", "該法", "本法", "前揭法", "前述法",
)
LAW_REFERENCE_DICTIONARY = set(LAW_WHITELIST) | set(LAW_ALIASES) | {
    "刑法", "毒品危害防制條例", "洗錢防制法",
}
KNOWN_LAW_NAME_SUFFIXES = tuple(sorted(LAW_REFERENCE_DICTIONARY, key=len, reverse=True))
LAW_REFERENCE_RE = re.compile(
    rf"(?P<reference>{'|'.join(map(re.escape, LAW_REFERENCE_NAMES))})"
    rf"(?P<article>第[{NUM}]+條(?:之[{NUM}]+)?(?:第[{NUM}]+項)?(?:第[{NUM}]+款)?)"
)
ELIDED_LAW_ARTICLE_RE = re.compile(
    rf"(?P<connector>[、，及與暨]|準用)"
    rf"(?P<article>第[{NUM}]+條(?:之[{NUM}]+)?(?:第[{NUM}]+項)?(?:第[{NUM}]+款)?)"
)
EXPLICIT_LAW_RE = re.compile(
    rf"(?P<name>[^，。；、：: （）()\s]{{1,30}}?(?:條例|法))(?=第[{NUM}]+條)"
)
# 罰金金額允許千分位逗號，但不併入 NUM，避免法條條號誤接受逗號。
TOTAL_RE = re.compile(
    rf"應執(?:行|刑)(?:之刑為)?(死刑|無期徒刑|有期徒(?:刑)?[{NUM}]+年(?:[{NUM}]+月)?|"
    rf"有期徒(?:刑)?[{NUM}]+月|拘役[{NUM}]+日|罰金(?:新臺幣)?[{NUM},]+元)"
)
# 主刑與罰金分欄時，主刑只保留死刑、無期、有期與拘役。
TOTAL_PRIMARY_RE = re.compile(
    rf"應執(?:行|刑)(?:之刑為)?(死刑|無期徒刑|有期徒(?:刑)?[{NUM}]+年(?:[{NUM}]+月)?|"
    rf"有期徒(?:刑)?[{NUM}]+月|拘役[{NUM}]+日)"
)
TOTAL_FINE_RE = re.compile(rf"應執(?:行|刑)(?:之刑為)?(罰金(?:新臺幣)?[{NUM},]+元)")
SINGLE_RE = re.compile(
    rf"(?<!易)(?:處|科)(死刑|無期徒刑|有期徒(?:刑)?[{NUM}]+年(?:[{NUM}]+月)?|"
    rf"有期徒(?:刑)?[{NUM}]+月|拘役[{NUM}]+日|罰金(?:新臺幣)?[{NUM},]+元)"
)
PRIMARY_RE = re.compile(
    rf"(?<!易)(?:處|科)(死刑|無期徒刑|有期徒(?:刑)?[{NUM}]+年(?:[{NUM}]+月)?|"
    rf"有期徒(?:刑)?[{NUM}]+月|拘役[{NUM}]+日)"
)
FINE_RE = re.compile(rf"(?<!易)(?:處|科|併科)(罰金(?:新臺幣)?[{NUM},]+元)")
CN_DIGITS = {
    "零": 0, "〇": 0, "○": 0, "一": 1, "壹": 1, "二": 2, "兩": 2, "貳": 2,
    "三": 3, "參": 3, "叁": 3, "叄": 3, "四": 4, "肆": 4, "五": 5, "伍": 5,
    "六": 6, "陸": 6, "七": 7, "柒": 7, "八": 8, "捌": 8,
    "九": 9, "玖": 9,
}
CN_UNITS = {"十": 10, "拾": 10, "百": 100, "佰": 100, "千": 1000, "仟": 1000,
            "萬": 10_000, "億": 100_000_000}


# 移除排版空白，讓跨法院、跨年度的相同句型可共用規則。
def compact(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def split_defendant_names(value) -> list[str]:
    """分號分隔姓名紀錄；空格、間隔號及頓號均保留，不推斷身分或去重。"""
    return [name.strip() for name in re.split(r"[；;]", str(value or "")) if name.strip()]


# 截掉誤併入主文欄的事實與理由，避免把後文論罪或前科當成本案主文。
def main_section(verdict: str) -> str:
    marker = re.search(r"主文[：:]", verdict)
    if marker:
        verdict = verdict[marker.end():]
    return MAIN_END_RE.split(verdict, maxsplit=1)[0]


# 將中文大小寫數字或阿拉伯數字轉為整數，供刑期月份正規化使用。
# 支援阿拉伯數字與中文單位混寫（如「10萬」「3萬6千」）：逐字掃描時，連續的
# 阿拉伯數字要當成一個多位數整體讀入，不能像中文數字那樣一字一位，否則
# 「10萬」會因為阿拉伯字元不在 CN_DIGITS 裡被忽略、算成 0。
def chinese_number(value: str) -> int:
    if value.isdigit():
        return int(value)
    total = section = number = 0
    i, length = 0, len(value)
    while i < length:
        char = value[i]
        if char.isdigit():
            j = i + 1
            while j < length and value[j].isdigit():
                j += 1
            number = int(value[i:j])
            i = j
            continue
        if char in CN_DIGITS:
            number = CN_DIGITS[char]
        elif char in CN_UNITS:
            unit = CN_UNITS[char]
            if unit >= 10_000:
                section = (section + number) * unit
                total += section
                section = number = 0
            else:
                section += (number or 1) * unit
                number = 0
        i += 1
    return total + section + number


# 保留原出現順序並去除重複標籤，避免同一法條在主文與理由重複輸出。
def unique(values):
    return list(dict.fromkeys(value for value in values if value))


# 在同一編號大點內將「同法／前開條例」回指至前方最近的完整法名。
def resolve_law_references(text: str, law_names) -> str:
    """只在 ㈠／（一）／一、等同一大點內解析承前指稱，不跨段推測。"""
    known_names = sorted(set(law_names), key=len, reverse=True)
    resolved_points = []
    # 補「第10條1項」的省字只用於匹配視圖，不修改來源文字。
    text = re.sub(rf"(第[{NUM}]+條(?:之[{NUM}]+)?)([{NUM}]+項)", r"\1第\2", compact(text))
    text = re.sub(rf"({'|'.join(map(re.escape, known_names))})(?=[{NUM}]+條)", r"\1第", text)
    text = re.sub(rf"((?:{'|'.join(map(re.escape, known_names))})第[{NUM}]+)(?=第[{NUM}]+項)", r"\1條", text)
    # 「犯罪事實一（二）」是句內回指，不是新的編號段落。
    starts = [0] + [m.start() for m in LAW_POINT_RE.finditer(text)
                    if m.start() and not re.search(rf'(?:犯罪事實|事實|附表)[{NUM}]*$', text[max(0, m.start()-16):m.start()])]
    for start, end in zip(starts, starts[1:] + [len(text)]):
        point = text[start:end]
        tokens = []
        for match in EXPLICIT_LAW_RE.finditer(point):
            raw_name = match.group("name")
            if any(raw_name.endswith(reference) for reference in LAW_REFERENCE_NAMES):
                continue
            recognized = next((name for name in known_names if raw_name.endswith(name)), "")
            tokens.append((match.start(), "law", recognized, raw_name, match))
        for match in LAW_REFERENCE_RE.finditer(point):
            tokens.append((match.start(), "reference", match.group("reference"), "", match))
        for match in ELIDED_LAW_ARTICLE_RE.finditer(point):
            tokens.append((match.start(), "elided", match.group("article"), "", match))
        for match in re.finditer(rf'同條(?P<unit>第[{NUM}]+項(?:第[{NUM}]+款)?)', point):
            tokens.append((match.start(), 'same_article', match['unit'], '', match))
        # 「甲罪、第2項之乙罪」明示承接項次；限同段、緊接罪名的列舉。
        for match in re.finditer(rf'(?<=罪)[、及與暨](?P<unit>第[{NUM}]+項(?:第[{NUM}]+款)?)(?=之)', point):
            tokens.append((match.start(), 'same_article', match[0][0] + match['unit'], '', match))
        tokens.sort(key=lambda token: (token[0], token[1] == "reference"))

        previous_laws = []
        previous_article = ''
        replacements = []
        for _, kind, value, raw_name, match in tokens:
            if kind == "law":
                # 未識別的法名仍會阻斷向更遠處回指，避免跨過最近候選誤配。
                previous_laws.append((raw_name.endswith("條例"), value))
                article = re.match(rf'第[{NUM}]+條(?:之[{NUM}]+)?', point[match.end():])
                previous_article = value + article[0] if value and article else ''
                continue
            if kind == 'same_article':
                if previous_article:
                    connector = value[0] if value[0] in '、及與暨' else ''
                    replacements.append((match.start(), match.end(), connector + previous_article + value[len(connector):]))
                continue
            if kind == "elided":
                referent = next((name for _, name in reversed(previous_laws) if name), "")
                if referent:
                    replacements.append((
                        match.start(), match.end(), match.group("connector") + referent + value,
                    ))
                    previous_laws.append((referent.endswith("條例"), referent))
                    previous_article = LAW_ID_RE.match(referent + value)[0]
                continue
            wants_regulation = value.endswith("條例")
            referent = next(
                (name for is_regulation, name in reversed(previous_laws)
                 if is_regulation == wants_regulation),
                "",
            )
            if referent:
                replacements.append((match.start(), match.end(), referent + match.group("article")))
                previous_laws.append((wants_regulation, referent))
                previous_article = LAW_ID_RE.match(referent + match.group('article'))[0]
            else:
                previous_article = ''
        for start, end, replacement in reversed(replacements):
            point = point[:start] + replacement + point[end:]
        resolved_points.append(expand_law_subunits(point))
    return "".join(resolved_points)


def expand_law_subunits(text: str) -> str:
    """只展開緊接完整法條的連續項／款，不跨罪名或敘述承接。"""
    replacements = []
    for match in LAW_RE.finditer(text):
        previous = clean_law(match)
        end = match.end()
        while True:
            tail = re.match(rf"[、及與暨](第[{NUM}]+(?P<level>項|款)(?:第[{NUM}]+款)?)", text[end:])
            if not tail:
                break
            base = re.sub(rf"第[{NUM}]+項.*$", "", previous) if tail['level'] == '項' else re.sub(rf"第[{NUM}]+款$", "", previous)
            previous = base + tail.group(1)
            replacements.append((end, end + tail.end(), '、' + previous))
            end += tail.end()
    for start, end, value in reversed(replacements):
        text = text[:start] + value + text[end:]
    return text


# 以使用者指定的 [A,B,C] 文字格式寫入多標籤欄位。
def labels(values) -> str:
    return "[" + ",".join(unique(values)) + "]"


# 刑期清單保留重複值，因為相同刑度可能分別屬於不同罪或不同被告。
def sentence_labels(values) -> str:
    return "[" + ",".join(values) + "]"


# 依主文是否同時包含定罪與非有罪諭知，判斷整案結果。
def case_result(verdict: str) -> str:
    verdict = main_section(verdict)
    non_guilty = [word for word in NON_GUILTY_WORDS if word in verdict]
    guilty = bool(re.search(
        r"犯.{0,80}?罪|處(?:死刑|無期徒刑|有期徒刑|拘役|罰金)|(?<!易)科罰金",
        verdict,
    ))
    if guilty and non_guilty:
        return "混合"
    if guilty:
        return "有罪"
    return "／".join(non_guilty) if non_guilty else "其他"


# 清掉法條前被正則一併捕捉的論述文字，只留下法規名稱與條項。
def clean_law(match: re.Match) -> str:
    name = match.group(1)
    for marker in (
        "分別係犯", "分別犯", "係幫助犯", "係共同犯", "係犯", "共同犯", "違反", "適用", "應依", "依",
        "該罪為", "無論以", "均係", "係",
    ):
        name = name.rsplit(marker, 1)[-1]
    name = re.sub(r"^(?:附錄本案所犯法條全文)?(?:中華民國)?", "", name)
    name = re.sub(r'^(?:修正前|修正後)', '', name)
    name = next((known for known in KNOWN_LAW_NAME_SUFFIXES if name.endswith(known)), name)
    return name + match.group(0).split(match.group(1), 1)[1]


# 同一法條的全文標題、主條與詳細條項只留最精確版本；不同項款不互相覆蓋。
def specific_substantive_laws(laws: list[str]) -> list[str]:
    """排除程序／施行法，並將同一法條的簡略重複值收斂為最詳細引文。"""
    kept: list[str] = []
    for law in laws:
        if any(excluded in law for excluded in ("刑事訴訟法", "刑法施行法")):
            continue
        match = LAW_ID_RE.match(law)
        identity = match.groups() if match else None
        replaced = False
        for index, existing in enumerate(kept):
            existing_match = LAW_ID_RE.match(existing)
            if not identity or not existing_match or existing_match.groups() != identity:
                continue
            if existing == law or existing.startswith(law):
                replaced = True
                break
            if law.startswith(existing):
                kept[index] = law
                replaced = True
                break
        if not replaced:
            kept.append(law)
    return kept


# 依理由、事實、適用法條的順序找法條，先命中即停止，避免較弱來源混入。
def extract_laws(reason: str, applicable_laws: str, facts: str = "") -> list[str]:
    found = []
    for text, conclusions_only in ((reason, True), (facts, True), (applicable_laws, False)):
        passages = [
            sentence for sentence in re.split(r"[。；]", conviction_context(text))
            if not conclusions_only or CONCLUSION_CHARGE_RE.search(sentence)
            or ("違反" in sentence and "罪" in sentence)
        ]
        found = specific_substantive_laws([
            clean_law(match) for passage in passages for match in LAW_RE.finditer(passage)
        ])
        if found:
            break
    return unique(found)


# 主文優先接受明載「犯○○罪」的結果，並移除罪數與累犯尾碼。
def extract_verdict_charges(verdict: str) -> list[str]:
    charges = VERDICT_CHARGE_RE.findall(main_section(verdict))
    return unique(
        CHARGE_COUNT_SUFFIX_RE.sub("", charge).removeprefix("共同") for charge in charges
        if not APPENDIX_REFERENCE_RE.match(charge)
        and not any(marker in charge for marker in ("，處", "沒收", "罪所得新臺幣", "無罪", "免訴", "不受理"))
    )


# 主文未寫「犯」或「罪」時，以被告名後、宣告刑前的文字作為罪名描述。
# 併科、易科與應執行不是單罪宣告刑起點；附表引用與程序性主文也不收。
def _first_declared_sentence(clause):
    return min((match for match in (
        PRIMARY_RE.search(clause),
        re.search(rf"(?<![易併])(?:處|科)罰金(?:新臺幣)?[{NUM},]+元", clause),
    ) if match), key=lambda match: match.start(), default=None)


def extract_verdict_descriptions(verdict: str, defendants: str = "") -> list[str]:
    descriptions = []
    names = sorted((compact(name) for name in split_defendant_names(defendants)), key=len, reverse=True)
    for clause in re.split(r"[。；]", main_section(compact(verdict))):
        if re.match(r"\s*應執(?:行|刑)", clause):
            continue
        sentence = _first_declared_sentence(clause)
        if not sentence:
            continue
        description = clause[:sentence.start()].rstrip("，,均各")
        description = description.removeprefix("又").removeprefix("另").removeprefix("被告")
        changed = True
        while changed:
            changed = False
            for name in names:
                if description.startswith(name):
                    description = description[len(name):].lstrip("、，,")
                    changed = True
        description = re.sub(rf"，(?:均)?累犯(?:，共?[{NUM}]+罪)?$", "", description)
        description = re.sub(rf"，共?[{NUM}]+罪$", "", description)
        if (len(description) >= 2
                and not APPENDIX_REFERENCE_RE.match(description)
                and not any(word in description for word in ("無罪", "免訴", "不受理", "撤銷", "駁回"))):
            descriptions.append(description)
    return unique(descriptions)


def verdict_clause_charges(clause: str, defendants: str = "") -> list[str]:
    """只在主文子句的受判決者位置依序試實名、遮名／代號、既有規則。"""
    text = compact(clause)
    subject = re.sub(rf"^(?:[{NUM}]+[、．.])?主刑部分[：:]?", "", text)
    subject = re.sub(rf"^(?:[{NUM}]+[、．.]|[（(][{NUM}]+[）)])", "", subject)
    names = sorted((compact(name) for name in split_defendant_names(defendants)), key=len, reverse=True)
    actual = "|".join(re.escape(name) for name in names)
    anonymous = r"[甲乙丙丁戊己庚辛壬癸][○〇◯ＯO]{2,}|[A-ZＡ-Ｚ][0-9０-９]{2,}"
    for pattern in (actual, anonymous):
        if not pattern:
            continue
        match = re.match(rf"^(?:被告)?(?:{pattern})(?:[、，,及與](?:{pattern}))*", subject)
        if not match:
            continue
        tail = subject[match.end():].lstrip("，,")
        # 本輪不改附表引用的既有解析／准入行為。
        sentence = _first_declared_sentence(tail)
        if "附表" in (tail[:sentence.start()] if sentence else tail):
            break
        charges = extract_verdict_charges(tail)
        if not charges:
            charges = [re.sub(r"^(?:均|各)?(?:又|另)?(?:共同)?犯", "", charge)
                       for charge in extract_verdict_descriptions(tail)]
        return charges
    return extract_verdict_charges(text) or extract_verdict_descriptions(text, defendants)


def conviction_context(text: str) -> str:
    """只遮蔽局部比較試算與明示量刑引用，不刪整句或依條號猜測。"""
    # 比較中的「修正後係犯」不是採用結論；後續論罪仍須原有逐罪配對。
    text = VERSION_COMPARISON_CLAUSE_RE.sub('【新舊法比較】', text)
    for match in reversed(list(LAW_RE.finditer(text))):
        if SENTENCING_CITATION_TAIL_RE.match(text[match.end():]):
            # 保留敘述邊界，避免刪掉量刑引用後把前罪法條接到下一罪。
            text = text[:match.start(1)] + '【量刑引用】' + text[match.end():]
    return text


# 輔助文字從已命中的法條結尾定位罪名，避免罪名內的「之」造成截斷。
def conviction_passages(text: str, conclusions_only=True):
    """保留法院論罪子句；檢方主張後須明示回到法院判斷才重新接受。"""
    for sentence in re.split(r"[。；]", conviction_context(text)):
        pending = sentence
        while pending:
            claim = PROSECUTION_CLAIM_RE.search(pending)
            passage = pending[:claim.start()] if claim else pending
            conclusion = CONCLUSION_CHARGE_RE.search(passage)
            if ((conclusion or not conclusions_only)
                    and not REJECTED_CONVICTION_RE.search(passage)):
                yield passage[conclusion.start():] if conclusion else passage
            if not claim:
                break
            resume = COURT_RESUMPTION_RE.search(pending, claim.end())
            pending = pending[resume.start():] if resume else ''


def extract_supporting_charges(text: str) -> list[str]:
    charges = []
    for sentence in conviction_passages(text):
        charges.extend(charge for _, charge in law_charge_pairs(sentence))
    return unique(charge.removeprefix("共同") for charge in charges if charge)


# 罪名以主文為準；主文沒有可用描述時，才依序退回理由與事實欄位。
def extract_charges(verdict: str, reason: str, facts: str = "", defendants: str = "") -> list[str]:
    return (unique(offense['charge'] for offense in extract_offense_clauses(verdict, defendants))
            or extract_supporting_charges(reason)
            or extract_supporting_charges(facts))


# 從單一文字片段取出實體法條，排除程序法與施行法。
def laws_in_text(text: str) -> list[str]:
    """保留片段內明載法條的出現順序。"""
    return specific_substantive_laws([clean_law(match) for match in LAW_RE.finditer(conviction_context(text))])


def charge_name_key(charge: str) -> str:
    """統一犯／罪尾綴寫法，保留罪名本身及未遂等區別。"""
    return re.sub(r"(?:之)?罪(?:嫌)?$", "", compact(charge).replace("共同犯", "共同")
                  .removeprefix("共同").removeprefix("幫助犯"))


def charge_matching_key(charge: str) -> str:
    """僅供逐罪配對的等義寫法；不覆寫主文罪名或刪除任意標點。"""
    key = charge_name_key(charge).replace('致人於死', '致死')
    key = re.sub(r'過失(?:傷害)?致(?:人)?重傷', '過失致重傷', key)
    key = re.sub(r'^未經許可，', '未經許可', key)
    key = re.sub(rf'(施用|持有)([{NUM}]+級毒品)', r'\1第\2', key)
    key = re.sub(rf'第([{NUM}]+)級',
                 lambda m: '第' + '一二三四'[chinese_number(m[1])-1] + '級'
                 if 1 <= chinese_number(m[1]) <= 4 else m[0], key)
    key = re.sub(rf'純質淨重([{NUM}]+)公克以上',
                 lambda m: f'純質淨重{chinese_number(m[1])}公克以上', key)
    key = re.sub(r'^(失火燒[燬毀])住宅、建築物等以外之物$', r'\1住宅等以外之物', key)
    key = re.sub(r'^明知為(禁藥|偽藥)而轉讓$', r'轉讓\1', key)
    key = re.sub(r'^意圖為自己(?:或第三人)?不法之所有，而侵占遺失物$', '侵占遺失物', key)
    return key


def charge_matches_law(charge: str, candidate: str, law: str) -> bool:
    if charge_matching_key(charge) == charge_matching_key(candidate):
        return True
    # 負扶助義務的描述不能配成一般遺棄，並保留致死、傷害等結果差異。
    return bool(re.fullmatch(r'對於無自救力之人，依(?:契約|法令|法令或契約)應扶助及保護而遺棄之',
                             charge_name_key(charge))
                and charge_name_key(candidate) == '遺棄'
                and re.fullmatch(r'刑法第294條第1項', law))


def contextual_charge_labels(charge, reason, facts, normalizer, allow_context_mixed=True):
    """先用罪名；僅不明駕駛罪讀取法院上下文，多駕駛罪須精確配對。"""
    standard = normalizer(charge)
    if not is_driving_charge(charge):
        return [standard] if standard else []
    explicit = driving_charge_kinds(charge)
    if explicit:
        if len(explicit) == 2 and not explicitly_mixed_driving(charge):
            return []
        # 主文明載兩種行為或兩款時保留兩種，不複製原始罪刑紀錄。
        return [kind for kind in ('酒駕罪', '毒駕罪') if kind in explicit]
    # 所有入口共用來源與主張篩選；明確罪名不必解析上下文。
    contexts = []
    for text in (reason, facts):
        court = split_prosecution_attachment(compact(text))[0]
        court = strip_law_quotations(resolve_law_references(court, LAW_REFERENCE_DICTIONARY))
        contexts.append('。'.join(conviction_passages(court, conclusions_only=False)))
    reason, facts = contexts
    kinds = set()
    same_offense_mixed = False
    conflicting_pair = False
    unrelated_passages = set()
    def outcome(name):
        return bool(re.search(r'致(?:人)?(?:於)?死|重傷', name))

    for text in (reason, facts):
        for passage in conviction_passages(text):
            if '否認' in passage:
                continue
            for laws, candidate in law_charge_groups(passage, charge_hint=charge):
                if not is_driving_charge(candidate):
                    continue
                if (outcome(charge) != outcome(candidate)
                        or ('重傷' in charge) != ('重傷' in candidate)):
                    unrelated_passages.add(passage)
                    continue
                if not allow_context_mixed and charge_matching_key(charge) != charge_matching_key(candidate):
                    unrelated_passages.add(passage)
                    continue
                name_kinds = driving_charge_kinds(candidate)
                law_kinds = set().union(*(driving_charge_kinds(law + '之不能安全駕駛罪') for law in laws))
                conflicting_pair |= bool(name_kinds and law_kinds and not name_kinds & law_kinds)
                kinds.update(name_kinds | law_kinds)
                # 兩個不同句子分別提到酒與毒，不視為同一罪雙重認定。
                same_offense_mixed |= explicitly_mixed_driving(candidate) or len(law_kinds) == 2
    # 配對已判為另一罪的句子，不能再從事實分支繞回採用。
    for passage in unrelated_passages:
        reason, facts = reason.replace(passage, ''), facts.replace(passage, '')
    if allow_context_mixed and mixed_driving_in_conviction(reason):
        kinds.update(('酒駕罪', '毒駕罪'))
        same_offense_mixed = True
    if conflicting_pair:
        return []
    if kinds:
        if len(kinds) == 1 or (same_offense_mixed and allow_context_mixed):
            return [kind for kind in ('酒駕罪', '毒駕罪') if kind in kinds]
        return []
    # 同案多個駕駛罪，未配對的行為敘述不足以定位該罪。
    if not allow_context_mixed:
        return []
    kinds = driving_fact_kinds(facts) | driving_fact_kinds(reason)
    # 單純事實層級同時出現酒毒仍不足以宣告雙標籤，須有主文／論罪認定。
    return list(kinds) if len(kinds) == 1 else []


def contextual_charge(charge, reason, facts, normalizer):
    """相容需要單一類別的呼叫；混合罪名由 contextual_charge_labels 提供。"""
    labels = contextual_charge_labels(charge, reason, facts, normalizer)
    return labels[0] if len(labels) == 1 else ""


def split_prosecution_attachment(text: str):
    """分離可辨認的附件正文；沒有附件標題時僅接受檢方結尾作來源證據。"""
    header = re.search(rf'附件[：:]?(?=臺灣.{{0,24}}檢察署檢察官(?:聲請簡易判決處刑書|起訴書)(?:[{NUM}]+年度|被告))', text)
    if header:
        return text[:header.start()], text[header.end():]
    if (re.search(r'依刑事訴訟法第(?:251|451|二百五十一|四百五十一)條.{0,35}(?:提起公訴|聲請逕以簡易判決處刑)', text)
            and not re.search(r'(?:判決|處刑|量處|判決處)如主文|應適用之法條[：:]', text)):
        return '', text
    return text, ''


def separate_source_fields(fields):
    """先逐欄辨別來源，再串接法院文字，避免附件結尾污染其他法院欄位。"""
    court_parts, attachments = [], []
    for label, value in fields:
        court, attachment = split_prosecution_attachment(compact(value))
        court_parts.append(court)
        if attachment:
            attachments.append((label, attachment))
    return '。'.join(part for part in court_parts if part), attachments


def law_citation_key(law: str):
    """比對引用時統一中文數字，輸出仍保留來源所載的法條。"""
    return re.sub(rf'[{NUM}]+', lambda m: str(chinese_number(m[0])), law)


def law_matching_sources(reason, facts, applicable_laws, extra_attachments=()):
    sources, attachments = [], list(extra_attachments)
    for label, text, conclusions_only in (
        ('理由', reason, True), ('事實', facts, True), ('適用法條', applicable_laws, False),
    ):
        court, attachment = split_prosecution_attachment(text)
        sources.append((label, court, conclusions_only, None))
        if attachment:
            attachments.append((label, attachment))
    court_text = '。'.join(text for label, text, _, _ in sources if label != '適用法條')
    adoption_passages = [p for p in conviction_passages(court_text, False)
                         if re.search(r'引用(?:(?:如)?附件(?:所示)?(?:之)?)?(?:臺灣.{0,24}檢察署)?(?:檢察官)?(?:聲?請(?:簡)?易判決處刑書|起訴書)|引用(?:如)?附件(?:所示)?之記載', p)]
    adopted_law = False
    for passage in adoption_passages:
        prefix = passage.split('引用', 1)[0]
        if not re.search(r'(?:犯罪)?事實.{0,45}(?:法條|論罪)', prefix):
            continue
        exception = re.search(r'除(.+)外', prefix)
        if exception and (laws_in_text(exception[1])
                          or re.search(r'(?:法條|論罪).{0,12}(?:更正|變更|刪除)', exception[1])):
            continue
        adopted_law = True
    # 明列法條更正／爭執時，附件仍須逐條得到法院結論佐證。
    if PROSECUTION_CLAIM_RE.search(court_text) or re.search(r'變更(?:起訴|聲請)?法條', court_text):
        adopted_law = False
    # 僅引用事實、證據不等於採用全部檢方論罪。
    conclusions = [p for p in conviction_passages(court_text, False)
                   if re.search(r'(?:判決|處刑|量處|判決處)如主文|應適用之法條[：:]', p)]
    allowed = {law_citation_key(law) for p in conclusions for law in laws_in_text(p)}
    if adoption_passages and (allowed or adopted_law):
        # 空集合表示法院明示採用法條；None 僅用於前面的法院來源。
        sources.extend((f'法院引用附件（{label}）', text, False, set() if adopted_law else allowed)
                       for label, text in attachments)
    return sources


# 先取得主文罪名，再從論罪段落取得該罪的法條；不使用案件層級法條補值。
def matched_laws(charge: str, standardized: str, main_clause: str,
                 reason: str, facts: str, applicable_laws: str,
                 normalizer=None, extra_attachments=(), diagnostics=None) -> tuple[list[str], str]:
    """名稱精確配對優先；同一粗類有多個不同罪名時，不猜測其對應。"""
    direct = laws_in_text(resolve_law_references(main_clause, LAW_REFERENCE_DICTIONARY))
    if direct:
        return direct, "主文同子句"
    candidates = []
    if diagnostics is not None:
        diagnostics['candidates'] = candidates
        diagnostics['incomplete_parallel_laws'] = []
    adopted_matches = []
    for source_name, text, conclusions_only, allowed in law_matching_sources(reason, facts, applicable_laws, extra_attachments):
        for passage in conviction_passages(text, conclusions_only):
            pairs = law_charge_pairs(passage, charge_hint=charge, diagnostics=diagnostics)
            exact = [(law, name) for law, name in pairs
                     if charge_matches_law(charge, name, law)
                     and (allowed is None or not allowed or law_citation_key(law) in allowed)]
            if allowed is not None:
                if exact:
                    adopted_matches.append((specific_substantive_laws(unique(law for law, _ in exact)), source_name))
                continue  # 附件不靠粗分類推測主文罪名。
            if exact:
                return specific_substantive_laws(unique(law for law, _ in exact)), source_name
            for law, name in pairs:
                identity = LAW_ID_RE.match(law)
                generic_driving = (charge_name_key(name) == '公共危險' and identity
                    and identity[1] == '刑法' and chinese_number(identity[2]) == 185
                    and chinese_number(identity[3] or '') == 3
                    and not re.search(r'致(?:人於)?死|重傷', charge))
                if is_driving_charge(charge) and (is_driving_charge(name) or generic_driving):
                    kind = driving_charge_kind(law + '之' + ('不能安全駕駛罪' if generic_driving else name))
                    if not standardized or not kind or kind == standardized:
                        candidates.append((law, name, source_name))
                elif normalizer and standardized and normalizer(name) == standardized:
                    # 保留既遂／未遂與加重結果差異，不能僅因粗類相同便掛錯法條。
                    if (('未遂' in charge) == ('未遂' in name)
                            and bool(re.search(r'致(?:人於)?死', charge)) == bool(re.search(r'致(?:人於)?死', name))
                            and ('重傷' in charge) == ('重傷' in name)
                            and ('電腦' in charge) == ('電腦' in name)
                            and re.findall(r'第[一二三四]級|純質淨重\d+公克以上', charge_matching_key(charge))
                            == re.findall(r'第[一二三四]級|純質淨重\d+公克以上', charge_matching_key(name))
                            and not (('取財' in charge and '得利' in name)
                                     or ('得利' in charge and '取財' in name))):
                        candidates.append((law, name, source_name))
    names = {charge_name_key(name) for _, name, _ in candidates}
    law_sets = {frozenset(law for law, name, _ in candidates if charge_name_key(name) == key)
                for key in names}
    # 多種罪名寫法只有在逐項法條集合完全相同時才不構成法條配對歧義。
    if len(names) == 1 or len(law_sets) == 1:
        source_name = candidates[0][2]
        return specific_substantive_laws(unique(law for law, _, source in candidates if source == source_name)), source_name
    if not candidates and adopted_matches:
        return adopted_matches[0]
    return [], ""


# 從已定位法條後取出罪名，供正式罪名與粗分類交叉比對。
def law_charge_pairs(text: str, charge_hint: str = "", diagnostics=None) -> list[tuple[str, str]]:
    return [(law, charge) for laws, charge in law_charge_groups(text, charge_hint, diagnostics) for law in laws]


def charge_hint_pattern(charge: str) -> str:
    """已知罪名才允許省略尾綴；等義詞限於固定詞界。"""
    name = charge_matching_key(charge)
    pattern = re.escape(name)
    pattern = pattern.replace('未經許可', '未經許可，?')
    pattern = pattern.replace('致死', '致(?:人於)?死')
    pattern = pattern.replace('住宅等以外之物', '住宅(?:、建築物)?等以外之物')
    for number, (chinese, formal) in enumerate(zip('一二三四', '壹貳參肆'), 1):
        pattern = pattern.replace(f'第{chinese}級', f'第?[{chinese}{formal}{number}]級')
    pattern = re.sub(r'純質淨重\d+公克以上', f'純質淨重[{NUM}]+公克以上', pattern)
    return pattern if len(name) >= 2 else ''


def parallel_drug_use_pairs(segment: str, laws):
    """明列兩級施用與第10條對應項次才分配；不以位置猜測、也不複製全部項次。"""
    match = re.match(r'之?施用第?([一二壹貳12])級(?:毒品)?[、及與暨](?:施用)?第?([一二壹貳12])級毒品(?:之?罪(?:嫌)?)?(?=[，、；。及與暨（）()]|$)', segment)
    if not match:
        return None
    grades = [chinese_number(match[1]), chinese_number(match[2])]
    assigned = {}
    for law in laws:
        unit = re.fullmatch(rf'毒品危害防制條例第([{NUM}]+)條第([{NUM}]+)項', law)
        if not unit or chinese_number(unit[1]) != 10 or chinese_number(unit[2]) not in grades:
            return []
        assigned[chinese_number(unit[2])] = law
    if len(set(grades)) != 2 or set(assigned) != set(grades):
        return []
    return [([assigned[grade]], f'施用第{"一二"[grade-1]}級毒品罪') for grade in grades]


def law_charge_groups(text: str, charge_hint: str = "", diagnostics=None):
    """把罪名前連續法條歸給該罪，並在下一個罪名後重新開始配對。"""
    pairs = []
    pending_laws = []
    # 省略「罪」只用已知主文罪名作錨點，不把法條後任意敘述當新罪名。
    name = charge_hint_pattern(charge_hint)
    hinted_charge = re.compile(
        rf"之?(?:共同)?({name})(?:之?罪(?:嫌)?)?(?=(?:之)?(?:犯行|行為)|之幫助犯|[，、；。及與暨（）()]|$)"
    ) if name else None
    law_matches = list(LAW_RE.finditer(text))
    for index, law_match in enumerate(law_matches):
        law = clean_law(law_match)
        if not law.startswith(("刑事訴訟法", "刑法施行法")):
            pending_laws.append(law)
        segment_end = len(text)
        if index + 1 < len(law_matches):
            following = law_matches[index + 1]
            # 法名正則可能吃進「共同犯…罪及」；罪名邊界須以清理後法名的位置為準。
            segment_end = following.start() + max(0, following.group(0).find(clean_law(following)))
        segment = text[law_match.end():segment_end].rstrip("及與暨")
        segment = re.sub(r"^(?:前段|後段)", "", segment)
        parallel = parallel_drug_use_pairs(segment, pending_laws)
        if parallel is not None:
            if not parallel and diagnostics is not None:
                diagnostics['incomplete_parallel_laws'].extend(pending_laws)
            # 展開只服務已知主文罪名，不憑理由新增附表未核實的罪名。
            if charge_hint:
                pairs.extend(parallel)
            pending_laws = []
            continue
        # 已知罪名先匹配，避免長罪名的前置解析強制要求「罪」。
        hinted = hinted_charge.match(segment) if hinted_charge else None
        if hinted and charge_matching_key(hinted[1]) == charge_matching_key(charge_hint):
            pairs.append((pending_laws, charge_hint))
            pending_laws = []
            continue
        charge_match = CHARGE_AFTER_LAW_RE.match(segment)
        if not charge_match and charge_hint:
            # 僅允許單一燒燬罪的標的列舉；其他罪之間的頓號仍是邊界。
            charge_match = re.match(
                r'之?((?:失火|放火)燒[燬毀](?:住宅|建築物)(?:、(?:住宅|建築物))*等以外之物罪)(?=[，、；。及與暨（）()]|$)', segment)
        if not charge_match:
            # 這些正式罪名的行為描述含逗號；不跨越其他罪名、法條或論述承接。
            charge_match = re.match(
                r"之?((?:汽車駕駛人|駕駛動力交通工具|未經許可)[^。；、（）()]{2,160}?罪)(?=[，、；。及與暨（）()]|$)",
                segment,
            )
        if not charge_match and law.startswith('刑法第185條之3'):
            # 酒精濃度＋服用毒品的單一罪名可能含頓號，仍保持同一組法條。
            driving_match = re.match(r"之?((?:吐氣所含酒精|服用酒類|服用毒品)[^。；]{2,160}?罪)(?=[，、；。及與暨（）()]|$)", segment)
            if driving_match and is_driving_charge(driving_match[1]):
                charge_match = driving_match
        if charge_match:
            charge = charge_match.group(1).removeprefix("共同")
            pairs.append((pending_laws, charge))
            pending_laws = []
        elif not re.fullmatch(r"[，、及與暨]*(?:準用)?", segment):
            # 遇到未解析的罪名／敘述就停止承接，避免把前罪法條掛到後罪。
            pending_laws = []
    return pairs


# 將法條後明載的罪名提供給粗分類比對，複合句已由 law_charge_pairs 分開。
def charges_after_laws(text: str) -> list[str]:
    """回傳同一片段內各法條所對應的明載罪名。"""
    return unique(charge for _, charge in law_charge_pairs(text))


# 以主文罪刑子句為單位，保留罪名、主刑與罰金的相同索引。
def extract_offense_clauses(verdict: str, defendants: str = "") -> list[dict[str, str]]:
    """將主文依出現順序拆成可對齊的罪刑紀錄。"""
    offenses = []
    for clause in re.split(r"[。；]|，(?=(?:又|另)(?!處|科|應執行))", main_section(compact(verdict))):
        charges = verdict_clause_charges(clause, defendants)
        if not charges:
            continue
        primaries = [normalize_sentence(value) for value in PRIMARY_RE.findall(clause)]
        fines = FINE_RE.findall(clause)
        size = max(len(charges), len(primaries), len(fines), 1)
        if len(charges) == 1 and size > 1:
            charges *= size
        for index, charge in enumerate(charges):
            offenses.append({
                "charge": charge,
                "primary": primaries[index] if index < len(primaries) else "",
                "fine": fines[index] if index < len(fines) else "",
                "clause": clause,
            })
    return offenses


# 讓舊工作表與精簡審核表共用罪刑對齊、承前法名及同罪名法條配對規則。
def extract_aligned_case(verdict: str, reason: str, facts: str, applicable_laws: str,
                         defendants: str, declared: list[str], normalizer, extra_attachments=(),
                         *, use_fallback=True):
    """回傳逐罪的主文罪刑、對應法條及來源；無主文罪名時才以理由備援。"""
    # 部分來源將「事實即理由」誤併入主文；僅重定位該明示段落，不掃全案補法條。
    embedded = re.search(r'事實即理由', verdict)
    if embedded:
        reason = verdict[embedded.end():] + '。' + reason
        verdict = verdict[:embedded.start()]
    reason = resolve_law_references(reason, LAW_REFERENCE_DICTIONARY)
    facts = resolve_law_references(facts, LAW_REFERENCE_DICTIONARY)
    applicable_laws = resolve_law_references(applicable_laws, LAW_REFERENCE_DICTIONARY)
    extra_attachments = [(label, resolve_law_references(text, LAW_REFERENCE_DICTIONARY))
                         for label, text in extra_attachments]
    offenses = extract_offense_clauses(verdict, defendants)
    charge_source = "主文" if offenses else ""
    court_reason = split_prosecution_attachment(reason)[0]
    court_facts = split_prosecution_attachment(facts)[0]
    supporting_charges = extract_supporting_charges(court_reason) or extract_supporting_charges(court_facts)
    if (not offenses and len(supporting_charges) == 1 and len(declared) == 1
            and "如附表" in verdict):
        offenses = [{
            "charge": supporting_charges[0], "primary": declared[0], "fine": "", "clause": verdict,
        }]
        charge_source = "主文刑期＋理由／事實正式罪名"
    if not offenses:
        fallback_charges = supporting_charges or extract_charges("", court_reason, court_facts, defendants)
        offenses = [{"charge": charge, "primary": "", "fine": "", "clause": ""}
                    for charge in fallback_charges]
        charge_source = "理由／事實備援" if offenses else ""

    laws_by_offense = []
    law_sources = []
    driving_count = sum(is_driving_charge(offense['charge']) for offense in offenses)
    for offense in offenses:
        charge = offense["charge"]
        offense["standard_charges"] = contextual_charge_labels(charge, court_reason, court_facts, normalizer, driving_count == 1)
        offense["standard_charge"] = offense["standard_charges"][0] if len(offense["standard_charges"]) == 1 else ""
        diagnostics = {}
        args = (
            charge, offense["standard_charge"], offense["clause"], reason, facts,
            applicable_laws, normalizer, extra_attachments,
        )
        laws, source_name = matched_laws(*args, diagnostics=diagnostics)
        offense['law_fallback'] = None
        if use_fallback and not laws:
            if charge_source != "主文":
                offense['law_fallback'] = dict(
                    charge=charge, legacy_raw=[], legacy_source='', evidence=[],
                    reasons=['no_main_verdict_offense'], recovered=False,
                )
            else:
                from law_fallback import try_fallback
                laws, audit = try_fallback(args, diagnostics)
                offense['law_fallback'] = dict(audit, recovered=bool(laws))
                if laws:
                    source_name = '舊版備援（' + audit['legacy_source'] + '）'
        laws_by_offense.append(laws)
        law_sources.append(source_name)
    return offenses, laws_by_offense, charge_source, law_sources


# 分開保留所有宣告主刑與明載的應執行主刑，罰金由專用欄位處理。
def extract_sentences(verdict: str) -> tuple[list[str], list[str]]:
    verdict = main_section(verdict)
    declared = [normalize_sentence(value) for value in PRIMARY_RE.findall(verdict)]
    executions = [normalize_sentence(value) for value in TOTAL_PRIMARY_RE.findall(verdict)]
    return declared, executions


# 應執行罰金屬案件層級，不硬配到單一罪名。
def extract_execution_fines(verdict: str) -> list[str]:
    """取出主文明載的應執行罰金。"""
    return TOTAL_FINE_RE.findall(main_section(verdict))


# 補回少數原文省略的「刑」字，使刑種與月份換算維持一致格式。
def normalize_sentence(sentence: str) -> str:
    return re.sub(r"^有期徒(?!刑)", "有期徒刑", sentence)


# 將有期徒刑換算成整月、拘役按 30 日一月換算至小數二位；其他刑種回傳 None。
# appendix_parser.py 也共用這支，避免主文與附表採用不同換算規則。
def sentence_to_months(sentence: str) -> Optional[Union[int, float]]:
    if sentence.startswith("拘役"):
        day = re.search(rf"([{NUM}]+)日", sentence)
        return round(chinese_number(day.group(1)) / 30, 2) if day else None
    if sentence.startswith("有期徒刑"):
        year = re.search(rf"([{NUM}]+)年", sentence)
        month = re.search(rf"([{NUM}]+)月", sentence)
        return (chinese_number(year.group(1)) * 12 if year else 0) + (
            chinese_number(month.group(1)) if month else 0)
    return None


# 將整案的有期徒刑與拘役清單換算成月，保留拘役的小數月數。
def sentence_months(sentences: list[str]) -> list[str]:
    return [f"{m:g}" for s in sentences if (m := sentence_to_months(s)) is not None]


# 提供案件層級的最高自由刑月份；死刑、無期徒刑與罰金不換算。
def maximum_sentence_months(sentences: list[str]) -> Union[int, float, str]:
    months = sentence_months(sentences)
    if not months:
        return ""
    maximum = max(map(float, months))
    return int(maximum) if maximum.is_integer() else maximum


# 區分法院明載執行刑、僅有宣告刑及缺少刑期，供模型篩選標籤品質。
def sentence_data_status(declared: list[str], executions: list[str], defendants: str) -> str:
    if executions:
        return "有應執行刑"
    if declared and len(split_defendant_names(defendants)) > 1 and len(declared) > 1:
        return "多被告未配對"
    if declared:
        return "僅宣告刑"
    return "無刑期"


# 結合主文、理由與適用法條，產生單一案件的一列結構化資料。
def structure_row(row: dict) -> list[object]:
    from normalize import normalize_charge

    verdict = compact(row.get("主文"))
    reason, reason_attachments = separate_source_fields((label, row.get(label)) for label in ('事實及理由', '理由'))
    facts, fact_attachments = separate_source_fields((label, row.get(label)) for label in ('事實', '犯罪事實'))
    applicable_laws = compact(row.get("適用法條"))
    defendants = row.get("被告／相對人") or ""
    declared, executions = extract_sentences(verdict)
    offenses, laws_by_offense, _, _ = extract_aligned_case(
        verdict, reason, facts, applicable_laws, defendants, declared, normalize_charge,
        reason_attachments + fact_attachments,
    )
    charges = [offense["charge"] for offense in offenses]
    laws = unique(law for offense_laws in laws_by_offense for law in offense_laws)
    result = case_result(verdict)
    sentence_status = sentence_data_status(declared, executions, defendants)
    status = "完整" if declared or executions else "無刑期"
    if ("如附表" in verdict and not charges) or (result in ("有罪", "混合") and (not laws or not charges)):
        status = "需人工核對"
    return [
        row.get("裁判字號", ""), row.get("裁判書連結", ""),
        row.get("裁判種類", ""), row.get("案件類型", ""), result,
        labels(laws), labels(charges), sentence_labels(declared), sentence_labels(sentence_months(declared)),
        maximum_sentence_months(declared), sentence_labels(executions),
        sentence_labels(sentence_months(executions)), bool(executions), sentence_status, status,
    ]


# 複製來源活頁簿並新增結構化工作表，保留原始資料供人工回查。
def structure_excel(source: Path, output: Path) -> int:
    if not source.exists():
        raise FileNotFoundError(f"找不到輸入檔：{source}")
    workbook = load_workbook(source)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"找不到工作表：{SOURCE_SHEET}")
    source_sheet = workbook[SOURCE_SHEET]
    headers = [cell.value for cell in source_sheet[1]]
    required = {
        "裁判字號", "裁判書連結", "裁判種類", "案件類型", "主文",
        "事實及理由", "理由", "適用法條",
    }
    missing = required - set(headers)
    if missing:
        raise ValueError(f"缺少必要欄位：{', '.join(sorted(missing))}")

    if OUTPUT_SHEET in workbook.sheetnames:
        del workbook[OUTPUT_SHEET]
    target = workbook.create_sheet(OUTPUT_SHEET)
    target.append([
        "裁判字號", "裁判書連結", "裁判種類", "案件類型", "案件結果",
        "法條", "罪名", "所有宣告刑", "宣告刑（月）", "宣告刑最大值（月）",
        "應執行刑", "應執行刑（月）", "是否明載應執行刑", "刑期資料狀態", "萃取狀態",
    ])
    for values in source_sheet.iter_rows(min_row=2, values_only=True):
        target.append(structure_row(dict(zip(headers, values))))

    for cell in target[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F3864")
    target.freeze_panes = "A2"
    target.auto_filter.ref = target.dimensions
    target.column_dimensions["A"].width = 38
    target.column_dimensions["B"].width = 55
    for column in "CDEFGHIJKLMNO":
        target.column_dimensions[column].width = 28
    for row in target.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output)
    return target.max_row - 1


# 執行最小自我檢查，確保中文數字、混合案件與總刑期規則未被改壞。
def self_check() -> None:
    assert chinese_number("貳仟壹佰捌拾捌萬肆仟肆佰") == 21_884_400
    assert case_result("甲犯竊盜罪，處有期徒刑參月。乙無罪。") == "混合"
    assert case_result("甲科罰金新臺幣壹萬元。") == "有罪"
    assert case_result("無罪。事實與理由：檢察官認被告犯罪。") == "無罪"
    declared, executions = extract_sentences(
        "各處有期徒刑參月；另處有期徒刑參月。應執行有期徒刑壹年貳月。"
    )
    assert declared == ["有期徒刑參月", "有期徒刑參月"]
    assert executions == ["有期徒刑壹年貳月"] and sentence_months(executions) == ["14"]
    assert extract_sentences("甲竊盜，科罰金新臺幣壹萬元。")[0] == []
    assert extract_offense_clauses("甲竊盜，科罰金新臺幣壹萬元。", "甲")[0]["fine"] == "罰金新臺幣壹萬元"
    assert extract_sentences("甲犯竊盜罪，處有期徒陸月。")[0] == ["有期徒刑陸月"]
    assert extract_sentences("應執刑有期徒刑肆年。")[1] == ["有期徒刑肆年"]
    assert extract_sentences(
        "處有期徒刑肆月，如易科罰金，以新臺幣壹仟元折算壹日。"
    )[0] == ["有期徒刑肆月"]
    assert extract_charges("", compact(
        "核被告所為，係犯刑法第185條之3第1項第3款之"
        "駕駛動力交通工具而有尿液所含毒品達行政院公告之品項及濃度值以上之情形罪。"
    )) == ["駕駛動力交通工具而有尿液所含毒品達行政院公告之品項及濃度值以上之情形罪"]
    assert extract_charges("", compact(
        "核被告所為，係犯毒品危害防制條例第10條第2項施用第二級毒品罪。"
    )) == ["施用第二級毒品罪"]
    assert extract_charges(
        compact("甲犯竊盜罪，處有期徒刑參月。"),
        compact("核被告所為，係犯刑法第335條第1項之侵占罪。"),
    ) == ["竊盜罪"]
    assert extract_charges(compact(
        "甲犯本判決附表乙編號一「罪名、宣告刑（含執行刑）」欄內所示之罪，"
        "處如附表所示之刑。"
    ), "") == []
    assert extract_charges(compact(
        "甲犯汽車駕駛人，行近行人穿越道不依規定讓行人優先通行，"
        "因而過失致人受傷罪，處有期徒刑參月。"
    ), "") == ["汽車駕駛人，行近行人穿越道不依規定讓行人優先通行，因而過失致人受傷罪"]
    assert extract_charges(compact(
        "甲犯竊盜罪，處有期徒刑參月。事實與理由：核被告所為，"
        "係犯刑法第335條第1項之侵占罪。"
    ), "") == ["竊盜罪"]
    assert structure_row({
        "主文": "甲犯如附表所示之罪，應執行有期徒刑壹年。",
    })[-1] == "需人工核對"


# 提供可直接套用預設檔名的命令列入口，也允許覆寫輸入與輸出路徑。
def main() -> None:
    parser = argparse.ArgumentParser(description="萃取每案罪名、法條、宣告刑與應執行刑")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.input.with_name(f"{args.input.stem}_structured.xlsx")
    self_check()
    count = structure_excel(args.input, output)
    print(f"完成：{count} 筆 → {output}")


if __name__ == "__main__":
    main()
