"""將罪法刑結構化結果另存為「標準化」工作表，提供常見法主條與罪名粗分類。"""

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from law_frequency import extract_law_keys
from driving_charge import is_driving_charge, driving_charge_kind


SOURCE_SHEET = "罪法刑結構化"
OUTPUT_SHEET = "標準化"
# 全資料原始引用前 200 名且至少 30 件的快照（150 類；排除指定三條）；詳見「白名單.md」。
# 跨時間驗證須只用訓練資料重建前 200 名並套用至少 30 件門檻，不可直接使用此快照。
LAW_WHITELIST = frozenset({
    "刑法第320條",
    "刑法第185條",
    "刑法第339條",
    "毒品危害防制條例第10條",
    "刑法第284條",
    "洗錢防制法第14條",
    "刑法第277條",
    "刑法第216條",
    "毒品危害防制條例第11條",
    "刑法第337條",
    "洗錢防制法第2條",
    "刑法第321條",
    "刑法第30條",
    "刑法第309條",
    "洗錢防制法第19條",
    "刑法第305條",
    "毒品危害防制條例第4條",
    "刑法第354條",
    "組織犯罪防制條例第3條",
    "刑法第268條",
    "毒品危害防制條例第2條",
    "刑法第336條",
    "刑法第304條",
    "刑法第135條",
    "刑法第276條",
    "刑法第231條",
    "刑法第140條",
    "商業會計法第71條",
    "商標法第97條",
    "刑法第266條",
    "刑法第310條",
    "刑法第335條",
    "家庭暴力防治法第61條",
    "性騷擾防治法第25條",
    "刑法第47條",
    "個人資料保護法第41條",
    "刑法第55條",
    "藥事法第83條",
    "毒品危害防制條例第20條",
    "洗錢防制法第16條",
    "刑法第306條",
    "刑法第302條",
    "刑法第214條",
    "洗錢防制法第3條",
    "公司法第9條",
    "刑法第315條",
    "銀行法第125條",
    "刑法第2條",
    "刑法第150條",
    "刑事訴訟法第310條",
    "刑法第201條",
    "刑事訴訟法第159條",
    "刑法第342條",
    "兒童及少年福利與權益保障法第112條",
    "刑法第210條",
    "刑法第224條",
    "妨害兵役治罪條例第3條",
    "刑法第225條",
    "家庭暴力防治法第2條",
    "刑法第227條",
    "刑法第1條",
    "槍砲彈藥刀械管制條例第8條",
    "刑法第346條",
    "毒品危害防制條例第23條",
    "刑法第217條",
    "刑法第169條",
    "刑法第221條",
    "槍砲彈藥刀械管制條例第7條",
    "證券交易法第171條",
    "刑事訴訟法第273條",
    "刑法第38條",
    "刑法第168條",
    "期貨交易法第112條",
    "刑法第319條",
    "毒品危害防制條例第5條",
    "刑法第57條",
    "刑事訴訟法第158條",
    "洗錢防制法第22條",
    "刑法第287條",
    "刑法第234條",
    "詐欺犯罪危害防制條例第44條",
    "刑法第59條",
    "刑法第175條",
    "刑法第11條",
    "毒品危害防制條例第8條",
    "刑事訴訟法第449條",
    "刑法第215條",
    "刑法第28條",
    "槍砲彈藥刀械管制條例第12條",
    "槍砲彈藥刀械管制條例第15條",
    "懲治走私條例第2條",
    "家庭暴力防治法第3條",
    "證券交易法第175條",
    "戶籍法第75條",
    "著作權法第91條",
    "著作權法第92條",
    "刑事訴訟法第451條",
    "兒童及少年性剝削防制條例第36條",
    "刑法第349條",
    "稅捐稽徵法第43條",
    "刑法第271條",
    "毒品危害防制條例第9條",
    "刑法第151條",
    "藥事法第82條",
    "槍砲彈藥刀械管制條例第14條",
    "醫療法第106條",
    "跟蹤騷擾防制法第18條",
    "刑法第13條",
    "刑法第220條",
    "刑法第62條",
    "毒品危害防制條例第17條",
    "刑法第138條",
    "刑法第171條",
    "刑法第25條",
    "刑事訴訟法第300條",
    "洗錢防制法第15條",
    "刑法第325條",
    "刑法第358條",
    "政府採購法第87條",
    "刑法第164條",
    "刑法第359條",
    "刑法第330條",
    "刑法第173條",
    "刑法第222條",
    "就業服務法第57條",
    "民法第74條",
    "毒品危害防制條例第24條",
    "廢棄物清理法第46條",
    "刑法第235條",
    "刑法第239條",
    "陸海空軍刑法第54條",
    "刑法第74條",
    "詐欺犯罪危害防制條例第43條",
    "性侵害犯罪防治法第21條",
    "臺灣地區與大陸地區人民關係條例第79條",
    "醫師法第28條",
    "刑法第35條",
    "個人資料保護法第20條",
    "稅捐稽徵法第41條",
    "刑事訴訟法第7條",
    "刑法第344條",
    "刑法第212條",
    "刑法第278條",
    "刑法第356條",
    "電子遊戲場業管理條例第22條",
    "刑法第255條",
    "性侵害犯罪防治法第50條",
    "貪污治罪條例第5條",
    "刑法第132條",
    "洗錢防制法第23條",
})
CHARGE_RULES = (
    (re.compile(r"尿液所含毒品|駕駛動力交通工具.*毒品|"
                r"(?:服用|施用|吸食).*毒品.*(?:不能安全駕駛|駕駛動力交通工具)"), "毒駕罪"),
    (re.compile(r"酒.*濃度|酒類.*駕駛"), "酒駕罪"),
    (re.compile(r"過失致重傷"), "過失致重傷罪"),
    (re.compile(r"過失傷害|過失.*(?:受傷|傷害)"), "過失傷害罪"),
    (re.compile(r"過失致死|過失致人於死"), "過失致死罪"),
    (re.compile(r"電腦.*不實財產權.*取財"), "電腦詐欺罪"),
    (re.compile(r"(?:自動付款|收費)設備"), "不正利用自動設備罪"),
    # 證券交易犯罪優先於一般詐欺；保留操縱價格與詐偽兩種細類。
    (re.compile(r"高買(?:低賣)?(?:有價)?證券|低賣(?:有價)?證券|"
                r"操縱(?:證券)?交易價格|操縱證券價格"), "操縱證券交易價格罪"),
    (re.compile(r"證券.*詐偽|詐偽.*證券"), "證券詐偽罪"),
    # 「詐術」須與財產行為相連，避免將施用毒品、性交等其他詐術行為當成詐欺。
    (re.compile(r"詐欺|詐術.*(?:交付|取財|得利|財產上不法.*利益)"), "詐欺罪"),
    (re.compile(r"竊盜"), "竊盜罪"),
    (re.compile(r"施用.*毒品"), "施用毒品罪"),
    (re.compile(r"持有.*毒品"), "持有毒品罪"),
    (re.compile(r"販賣.*毒品"), "販賣毒品罪"),
    (re.compile(r"轉讓.*毒品"), "轉讓毒品罪"),
    (re.compile(r"運輸.*毒品"), "運輸毒品罪"),
    (re.compile(r"製造.*毒品"), "製造毒品罪"),
    (re.compile(r"偽造.*有價證券|行使.*有價證券"), "偽造有價證券罪"),
    (re.compile(r"偽造署押"), "偽造署押罪"),
    (re.compile(r"填製不實.*(?:會計憑證|罪)|不實.*會計憑證"), "填製不實會計憑證罪"),
    (re.compile(r"行使.*文書|偽造.*文書|登載不實"), "偽造文書罪"),
    (re.compile(r"侵占.*"), "侵占罪"),
    (re.compile(r"竊佔"), "竊佔罪"),
    (re.compile(r"恐嚇"), "恐嚇罪"),
    (re.compile(r"毀損|毀棄損壞|致令他人物品不堪用"), "毀損罪"),
    (re.compile(r"肇事.*逃逸|交通事故.*逃逸"), "肇事逃逸罪"),
    (re.compile(r"重傷"), "重傷罪"),
    (re.compile(r"傷害"), "傷害罪"),
    (re.compile(r"洗錢"), "洗錢罪"),
    (re.compile(r"犯罪組織|參與組織"), "組織犯罪罪"),
    (re.compile(r"僭行公務員職權"), "僭行公務員職權罪"),
    (re.compile(r"侮辱公務員|公務員.*執行.*當場侮辱"), "侮辱公務員罪"),
    (re.compile(r"妨害公務|公務員.*施.*強暴|強暴侮辱"), "妨害公務罪"),
    (re.compile(r"公然侮辱|誹謗"), "妨害名譽罪"),
    (re.compile(r"強制性交"), "強制性交罪"),
    (re.compile(r"強制猥褻"), "強制猥褻罪"),
    (re.compile(r"殺人"), "殺人罪"),
    (re.compile(r"強盜"), "強盜罪"),
    (re.compile(r"搶奪"), "搶奪罪"),
    (re.compile(r"強制.*罪"), "強制罪"),
    (re.compile(r"放火"), "放火罪"),
    (re.compile(r"失火"), "失火罪"),
    (re.compile(r"公共危險罪"), "公共危險罪"),
    (re.compile(r"賭博|供給賭博場所|聚眾賭博"), "賭博罪"),
    (re.compile(r"違反保護令"), "違反保護令罪"),
    (re.compile(r"媒介.*性交|性交.*媒介|容留.*(?:性交|猥褻)"), "妨害風化罪"),
    (re.compile(r"未繳納股款|收回股款"), "未繳納股款罪"),
    (re.compile(r"背信"), "背信罪"),
    (re.compile(r"剝奪(?:他)?人.*行動自由|私行拘禁"), "剝奪行動自由罪"),
    (re.compile(r"侵害商標權"), "侵害商標權罪"),
    (re.compile(r"誣告"), "誣告罪"),
    (re.compile(r"性騷擾"), "性騷擾罪"),
    (re.compile(r"轉讓(?:禁藥|偽藥)"), "轉讓禁藥罪"),
    (re.compile(r"妨害兵役|妨害.*(?:徵兵|召集)|妨害役男"), "妨害兵役罪"),
    (re.compile(r"偽證"), "偽證罪"),
    (re.compile(r"侵入.*(?:住宅|建築物)"), "侵入住宅罪"),
    (re.compile(r"公然猥褻"), "公然猥褻罪"),
    (re.compile(r"乘機性交"), "乘機性交罪"),
    (re.compile(r"乘機猥褻"), "乘機猥褻罪"),
    (re.compile(r"非法經營證券業務"), "非法經營證券業務罪"),
    (re.compile(r"非法經營期貨"), "非法經營期貨業務罪"),
    (re.compile(r"內線交易"), "內線交易罪"),
    (re.compile(r"損壞公務員.*掌管"), "損壞公務物品罪"),
    (re.compile(r"損壞.*物品"), "損壞他人物品罪"),
    (re.compile(r"贓物"), "贓物罪"),
    (re.compile(r"竊錄|妨害秘密"), "妨害秘密罪"),
    (re.compile(r"(?:持有|寄藏|製造).*(?:槍枝|手槍|子彈|空氣槍|槍砲.*零件|爆裂物)"), "槍砲罪"),
    (re.compile(r"未經許可.*刀械|非法.*刀械"), "非法持有刀械罪"),
    (re.compile(r"妨害公眾往來安全"), "妨害公眾往來安全罪"),
    (re.compile(r"公共場所聚集三人以上|聚集三人以上.*強暴"), "聚眾施強暴罪"),
    (re.compile(r"非法經營銀行|非法.*收受存款|非法辦理.*匯兌"), "非法經營銀行業務罪"),
    (re.compile(r"私運管制物品"), "私運管制物品罪"),
    (re.compile(r"頂替"), "頂替罪"),
    (re.compile(r"重利"), "重利罪"),
    (re.compile(r"(?:相姦|通姦)"), "通姦罪"),
    (re.compile(r"(?:未滿十六歲|未滿16歲|未滿十四歲|未滿14歲).*(?:性交|猥褻)"), "與未成年人性交猥褻罪"),
    # 兒少性影像沿用兒少性剝削類別，優先於一般性影像規則。
    (re.compile(r"(?:少年|兒童|兒少|未滿(?:18|十八)歲).*性影像|"
                r"性影像.*(?:少年|兒童|兒少|未滿(?:18|十八)歲)"), "兒少性剝削罪"),
    (re.compile(r"性影像|妨害性隱私"), "妨害性隱私及不實性影像罪"),

    (re.compile(r"非法執行醫療"), "非法執行醫療業務罪"),
    (re.compile(r"輸入禁藥"), "輸入禁藥罪"),
    (re.compile(r"未經許可入國"), "未經許可入國罪"),
    (re.compile(r"損害債權"), "損害債權罪"),
    (re.compile(r"逃漏稅捐"), "逃漏稅捐罪"),
    (re.compile(r"侵害.*著作財產權|著作權法.*罪"), "侵害著作權罪"),
    (re.compile(r"非法營業"), "非法營業罪"),
    (re.compile(r"跟蹤騷擾"), "跟蹤騷擾罪"),
    (re.compile(r"商品虛偽標記"), "商品虛偽標記罪"),
    (re.compile(r"個人資料|個資法|非法利用個人資料"), "違反個人資料保護法罪"),
    (re.compile(r"性侵害犯罪防治法|加害人屆期不履行"), "違反性侵害防治義務罪"),
    (re.compile(r"盜用電信"), "盜用電信罪"),
    (re.compile(r"非法經營證券投資顧問"), "非法經營投資顧問業務罪"),
    (re.compile(r"使.*大陸地區人民非法進入"), "使大陸地區人民非法進入罪"),
    (re.compile(r"填載不實"), "填製不實會計憑證罪"),
    (re.compile(r"散布猥褻"), "散布猥褻物品罪"),
    (re.compile(r"少年.*猥褻|兒童.*猥褻"), "兒少性剝削罪"),
    (re.compile(r"違法重建"), "違法重建罪"),
    (re.compile(r"輸入私菸"), "輸入私菸罪"),
    (re.compile(r"非法(?:僱用|留用)"), "非法僱用罪"),
    (re.compile(r"無故.*(?:取得|變更|刪除).*電腦.*電磁紀錄"), "妨害電腦使用罪"),
    (re.compile(r"妨害醫事人員"), "妨害醫事人員執行業務罪"),
    (re.compile(r"偽造.*(?:紙幣|通用貨幣)"), "偽造貨幣罪"),
    (re.compile(r"重婚"), "重婚罪"),
    (re.compile(r"媒介.*猥褻|猥褻.*(?:媒介|容留)"), "妨害風化罪"),
    (re.compile(r"選舉.*(?:賄賂|交付)|交付賄賂"), "選舉賄賂罪"),
    (re.compile(r"發還股款"), "未繳納股款罪"),
    (re.compile(r"冒用身分.*身分證"), "冒用身分證罪"),
    (re.compile(r"使開標發生不正確結果"), "妨害投標罪"),
    (re.compile(r"受禁止出國處分而出國"), "違反禁止出國處分罪"),
    (re.compile(r"主管事務圖利|利用職務.*詐取財物"), "貪污罪"),
    (re.compile(r"財務報表.*不實"), "財報不實罪"),
    (re.compile(r"偽造印文"), "偽造印文罪"),
    (re.compile(r"護照交付他人|冒名使用.*護照"), "違反護照條例罪"),
    (re.compile(r"侵入.*(?:住居|住所)"), "侵入住宅罪"),
    (re.compile(r"湮滅.*證據"), "湮滅證據罪"),
    (re.compile(r"偽造.*(?:信用卡|支付卡)|行使偽造信用卡"), "偽造信用卡罪"),
    (re.compile(r"非法清理廢棄物"), "非法清理廢棄物罪"),
    (re.compile(r"散播.*疫情.*不實"), "散播疫情不實訊息罪"),
)


# 拆開 structure_tasks.py 使用的 [A,B,C] 多標籤格式，空清單不產生項目。
def split_labels(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text or text == "[]":
        return []
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1]
    return [item.strip() for item in text.split(",") if item.strip()]


# 將明載法條合併為主條號，再以「法律名稱＋主條號」逐條篩選。
def normalize_law_label(raw: object) -> str:
    """標準化單一法條；不推定「同法／該條例」的指涉。"""
    keys = extract_law_keys(raw)
    return keys[0] if keys and keys[0] in LAW_WHITELIST else ""


# 只保留逐條白名單內的明載名稱與主條號。
def normalize_law_levels(value: object) -> tuple[str, str]:
    normalized: list[str] = []
    for raw in split_labels(value):
        law = normalize_law_label(raw)
        if law and law not in normalized:
            normalized.append(law)
    if not normalized:
        return "", ""
    labels = "[" + ",".join(normalized) + "]"
    return labels, labels


# 以明確的犯罪家族規則標準化單一罪名；未遂只保留在原始罪名。
def normalize_charge(charge: object) -> str:
    """回傳不區分既遂、未遂的單一罪名粗分類；未命中時留空。"""
    text = str(charge or "").strip()
    if is_driving_charge(text):
        return driving_charge_kind(text)
    return next((label for pattern, label in CHARGE_RULES if pattern.search(text)), "")


# 舊版整格介面保留全有全無語意，對齊新欄位改用 normalize_charge。
def normalize_charges(value: object, reason: str = "", facts: str = "") -> str:
    from structure_tasks import contextual_charge_labels

    normalized: list[str] = []
    charges = split_labels(value)
    driving_count = sum(is_driving_charge(charge) for charge in charges)
    for charge in charges:
        standards = contextual_charge_labels(
            charge, reason, facts, normalize_charge, allow_context_mixed=driving_count == 1,
        )
        if not standards:
            return ""
        for standard in standards:
            if standard not in normalized:
                normalized.append(standard)
    return "[" + ",".join(normalized) + "]" if normalized else ""


# 讀取結構化工作表並新增標準化工作表；刑期維持原樣以免改變量刑語意。
def normalize_workbook(source: Path, output: Path) -> int:
    from structure_tasks import separate_source_fields

    if not source.exists():
        raise FileNotFoundError(f"找不到輸入檔：{source}")
    workbook = load_workbook(source)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"找不到工作表：{SOURCE_SHEET}")

    source_sheet = workbook[SOURCE_SHEET]
    headers = [cell.value for cell in source_sheet[1]]
    required = {"裁判字號", "法條", "罪名", "所有宣告刑", "應執行刑"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"缺少必要欄位：{', '.join(sorted(missing))}")

    if OUTPUT_SHEET in workbook.sheetnames:
        del workbook[OUTPUT_SHEET]
    target = workbook.create_sheet(OUTPUT_SHEET)
    target.append([
        "判決書名稱", "法條（原始）", "法條標準鍵", "法條粗分類",
        "罪名（原始）", "罪名（標準化）", "所有宣告刑", "應執行刑",
    ])
    # 舊工作表仍保留原始罪名；上下文僅從同案且唯一的原始資料列取回。
    source_rows = {}
    if "裁判書資料" in workbook.sheetnames:
        source_values = workbook["裁判書資料"].iter_rows(values_only=True)
        source_headers = next(source_values, ())
        for values in source_values:
            raw = dict(zip(source_headers, values))
            key = (raw.get("裁判字號"), raw.get("裁判書連結"))
            source_rows[key] = raw if key not in source_rows else None
    for values in source_sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        standard_laws, coarse_laws = normalize_law_levels(row.get("法條"))
        raw = source_rows.get((row.get("裁判字號"), row.get("裁判書連結"))) or {}
        reason, _ = separate_source_fields((key, raw.get(key)) for key in ("事實及理由", "理由"))
        facts, _ = separate_source_fields((key, raw.get(key)) for key in ("事實", "犯罪事實"))
        target.append([
            row.get("裁判字號", ""), row.get("法條", ""), standard_laws, coarse_laws,
            row.get("罪名", ""), normalize_charges(row.get("罪名"), reason, facts),
            row.get("所有宣告刑", ""), row.get("應執行刑", ""),
        ])

    for cell in target[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F3864")
    target.freeze_panes = "A2"
    target.auto_filter.ref = target.dimensions
    for column, width in zip("ABCDEFGH", (38, 45, 45, 45, 45, 35, 32, 32)):
        target.column_dimensions[column].width = width
    for row in target.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output)
    workbook.close()
    return target.max_row - 1


# 執行最小自我檢查，涵蓋逐條白名單、非開頭法名及常見罪名粗分類。
def self_check() -> None:
    assert normalize_law_levels("[刑法第一百八十五條之三第1項,同法第339條之4]") == (
        "[刑法第185條]", "[刑法第185條]"
    )
    assert normalize_law_levels(
        "[中華民國刑法第參佰貳拾條第1項,修正前刑法第320條]"
    ) == (
        "[刑法第320條]", "[刑法第320條]"
    )
    assert normalize_law_levels("[刑事訴訟法第158條]") == (
        "[刑事訴訟法第158條]", "[刑事訴訟法第158條]"
    )
    assert normalize_law_levels(
        "[㈱按甲基安非他命為毒品危害防制條例第2條第2項第2款,同條例第10條第2項]"
    ) == (
        "[毒品危害防制條例第2條]", "[毒品危害防制條例第2條]"
    )
    assert normalize_law_levels("[依陸海空軍刑法第54條第1項]") == (
        "[陸海空軍刑法第54條]", "[陸海空軍刑法第54條]"
    )
    assert normalize_law_levels("[前開商標法第95條第1項]") == ("", "")
    assert normalize_law_levels("[]") == ("", "")
    assert normalize_charges("[三人以上共同詐欺取財未遂罪,攜帶兇器竊盜罪]") == (
        "[詐欺罪,竊盜罪]"
    )
    assert normalize_charges("[吐氣所含酒精濃度達每公升0.25毫克以上而駕駛動力交通工具罪]") == (
        "[酒駕罪]"
    )
    assert normalize_charges("[駕駛動力交通工具而有尿液所含毒品達公告濃度值以上之情形罪]") == (
        "[毒駕罪]"
    )
    assert normalize_charges("[不能安全駕駛動力交通工具罪,圖利聚眾賭博罪]") == (
        ""
    )
    assert normalize_charges("[商標法第九十七條之非法販賣侵害商標權之商品罪]") == (
        "[侵害商標權罪]"
    )
    assert normalize_charges("[]") == ""
    assert normalize_charges("[尚未建立分類之罪]") == ""


# 提供原檔新增工作表的預設操作，也允許另存新檔以保留輸入檔。
def main() -> None:
    parser = argparse.ArgumentParser(description="新增罪法刑標準化工作表")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.input
    self_check()
    count = normalize_workbook(args.input, output)
    print(f"完成：{count} 筆 → {output}（{OUTPUT_SHEET}）")


if __name__ == "__main__":
    main()
