"""Convert the canonical structured XLSX sheet to SQLite tables and live filter views."""

import argparse
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


SOURCE_SHEET = "罪法刑"
TABLE = "罪法刑"
BATCH_SIZE = 1000


# Quote trusted workbook identifiers so Chinese names and punctuation remain valid SQL.
def quote(identifier: str) -> str:
    """Return a safely quoted SQLite identifier."""
    return '"' + identifier.replace('"', '""') + '"'


# Import the canonical main sheet in batches instead of loading every row into memory.
def import_main_sheet(source: Path, connection: sqlite3.Connection) -> int:
    """Replace the SQLite main table with rows from the canonical XLSX sheet."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"找不到工作表：{SOURCE_SHEET}")
    sheet = workbook[SOURCE_SHEET]
    headers = [str(cell.value) for cell in next(sheet.iter_rows(max_row=1))]
    if not headers or any(header in ("", "None") for header in headers):
        workbook.close()
        raise ValueError("主表標題列不完整")

    columns = ", ".join(f"{quote(header)} TEXT" for header in headers)
    connection.execute(f"DROP TABLE IF EXISTS {quote(TABLE)}")
    connection.execute(f"CREATE TABLE {quote(TABLE)} (id INTEGER PRIMARY KEY, {columns})")
    placeholders = ", ".join("?" for _ in headers)
    insert_sql = (
        f"INSERT INTO {quote(TABLE)} ({', '.join(map(quote, headers))}) "
        f"VALUES ({placeholders})"
    )
    count = 0
    batch = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        batch.append(tuple("" if value is None else str(value) for value in row))
        if len(batch) == BATCH_SIZE:
            connection.executemany(insert_sql, batch)
            count += len(batch)
            batch.clear()
    if batch:
        connection.executemany(insert_sql, batch)
        count += len(batch)
    workbook.close()
    return count


# Recreate worksheet-equivalent views from the main table so filters update together.
def create_views(connection: sqlite3.Connection) -> None:
    """Create live views matching the review worksheets without duplicating rows."""
    guilty = f"{quote('案件結果')} = '有罪'"
    denominator = f"{guilty} AND {quote('被告狀態')} = '單被告'"
    views = {
        "有罪案件": guilty,
        "有罪_單被告": denominator,
        "單一應執行刑": f"{denominator} AND {quote('單一刑期來源')} = '應執行刑'",
        "單一宣告刑": f"{denominator} AND {quote('單一刑期來源')} = '宣告刑'",
        "非單一刑期": f"{denominator} AND COALESCE({quote('單一刑期來源')}, '') = ''",
        "罪法_有值但標準化無值": f"{denominator} AND COALESCE({quote('標準化缺值欄位')}, '') <> ''",
        "罪法_原始無值": f"{denominator} AND COALESCE({quote('原始缺值欄位')}, '') <> ''",
        "刑期_有值但月數無值": f"{denominator} AND COALESCE({quote('刑期月數缺值欄位')}, '') <> ''",
        "刑期_原始無值": f"{denominator} AND COALESCE({quote('刑期原始缺值欄位')}, '') <> ''",
    }
    for name, where in views.items():
        connection.execute(f"DROP VIEW IF EXISTS {quote(name)}")
        connection.execute(
            f"CREATE VIEW {quote(name)} AS SELECT * FROM {quote(TABLE)} WHERE {where}"
        )
    connection.execute(
        f"CREATE INDEX IF NOT EXISTS idx_structured_result ON {quote(TABLE)} ({quote('案件結果')})"
    )
    connection.execute(
        f"CREATE INDEX IF NOT EXISTS idx_structured_sentence_status "
        f"ON {quote(TABLE)} ({quote('刑期資料狀態')})"
    )


# Provide one repeatable command for refreshing the browser database after extraction.
def main() -> None:
    """Parse CLI paths, replace the SQLite data, and report the imported row count."""
    parser = argparse.ArgumentParser(description="將罪法刑 XLSX 轉成 Datasette 可瀏覽的 SQLite")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    if not args.input.exists():
        parser.error(f"找不到輸入檔：{args.input}")
    output = args.output or args.input.with_suffix(".db")
    with sqlite3.connect(output) as connection:
        count = import_main_sheet(args.input, connection)
        create_views(connection)
    print(f"完成：{count} 筆 → {output}")


if __name__ == "__main__":
    main()
