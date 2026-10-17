"""Build the compact criminal-judgment review workbook from raw exported judgments."""

import argparse
import json
import os
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from normalize import normalize_charge, normalize_law_label
from structure_tasks import (
    case_result,
    compact,
    extract_aligned_case,
    extract_execution_fines,
    extract_sentences,
    sentence_months,
    sentence_to_months,
)


SOURCE_SHEET = "裁判書資料"
SHEETS = (
    "罪法刑", "有罪_單被告", "單一應執行刑", "單一宣告刑", "非單一刑期",
    "罪法_有值但標準化無值", "罪法_原始無值", "刑期_有值但月數無值", "刑期_原始無值",
)
HEADERS = (
    "裁判字號", "裁判書連結", "被告／相對人", "被告狀態", "裁判種類", "案件類型", "案件結果", "主文",
    "罪名（原始）", "罪名（標準化）", "罪名來源",
    "法條（對應原始）", "法條（對應標準鍵）", "法條來源",
    "宣告主刑", "宣告主刑（月）", "罰金", "應執行刑", "應執行刑（月）", "應執行罰金",
    "單一刑期", "單一刑期（月）", "單一刑期來源", "單一刑期風險",
    "刑期資料狀態", "萃取狀態", "原始缺值欄位", "標準化缺值欄位",
    "刑期月數缺值欄位", "刑期原始缺值欄位",
)
# Empty strings and serialized empty lists both represent missing extraction values.
def has_value(value: object) -> bool:
    """Return whether an exported field contains a usable value."""
    return value is not None and str(value).strip() not in {"", "[]"}


# Serialize aligned lists as valid JSON so commas inside charge names do not break positions.
def json_value(value: object) -> str:
    """輸出不轉義中文的緊湊 JSON 字串。"""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


# Require defendant-like names to appear in the main text; polluted party fragments stay unknown.
def classify_defendants(defendants: str, verdict: str) -> str:
    """Classify a case as single, multiple, or unknown using names evidenced in the main text."""
    names = {
        part for part in re.split(r"[；;]", compact(defendants))
        if re.fullmatch(r"[\u4e00-\u9fff○Ａ-ＺA-Za-z]{2,8}", part)
        and part in compact(verdict)
    }
    return "單被告" if len(names) == 1 else "多被告" if len(names) > 1 else "不明"


# Create a styled write-only sheet so large review files stay memory-efficient.
def add_sheet(workbook: Workbook, name: str):
    """Create one output sheet with the shared review header."""
    sheet = workbook.create_sheet(name)
    cells = []
    for value in HEADERS:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cells.append(cell)
    sheet.append(cells)
    sheet.freeze_panes = "A2"
    return sheet


# Convert one source row into the shared schema used by the main sheet and all review subsets.
def build_row(source: dict) -> tuple[list[object], dict[str, bool]]:
    """Run extraction and normalization once, returning the row plus filter flags."""
    verdict = compact(source.get("主文"))
    reason = compact(source.get("事實及理由")) + compact(source.get("理由"))
    facts = compact(source.get("事實")) + compact(source.get("犯罪事實"))
    applicable_laws = compact(source.get("適用法條"))
    defendants = compact(source.get("被告／相對人"))
    defendant_status = classify_defendants(defendants, verdict)
    declared, executions = extract_sentences(verdict)
    execution_fines = extract_execution_fines(verdict)
    offenses, raw_laws, charge_source, law_sources = extract_aligned_case(
        verdict, reason, facts, applicable_laws, defendants, declared, normalize_charge,
    )
    raw_charges = [offense["charge"] for offense in offenses]
    standard_charges = [normalize_charge(charge) for charge in raw_charges]
    standard_laws = [[normalize_law_label(law) for law in laws] for laws in raw_laws]
    primaries = [offense["primary"] for offense in offenses]
    primary_months = [sentence_to_months(sentence) or "" if sentence else "" for sentence in primaries]
    fines = [offense["fine"] for offense in offenses]
    result = case_result(verdict)

    # 單一刑期只能由輸出欄位衍生，不得另用全文刑期結果繞過罪刑對齊。
    if len(executions) == 1:
        selected, sentence_source = executions, "應執行刑"
    elif not executions and len(primaries) == 1 and primaries[0]:
        selected, sentence_source = primaries, "宣告刑"
    else:
        selected, sentence_source = [], ""
    risk = "多被告且單一應執行刑，可能只對應部分被告" if defendant_status == "多被告" and len(executions) == 1 else ""
    charge_missing = not raw_charges
    law_missing = not any(raw_laws)
    raw_missing = "、".join(label for label, condition in (("罪", charge_missing), ("法", law_missing)) if condition)
    standard_missing = "、".join(label for label, condition in (
        ("罪", bool(raw_charges) and any(not charge for charge in standard_charges)),
        ("法", any(raw_laws) and any(not standard for laws in standard_laws for standard in laws)),
    ) if condition)
    month_missing = "、".join(label for label, condition in (
        ("宣告刑", bool(declared) and not sentence_months(declared)),
        ("應執行刑", bool(executions) and not sentence_months(executions)),
    ) if condition)
    sentence_missing = "、".join(label for label, condition in (
        ("宣告刑", not declared), ("應執行刑", not executions),
    ) if condition)
    sentence_status = "有應執行刑" if executions else "有應執行罰金" if execution_fines else "僅宣告主刑" if declared else "僅罰金" if any(fines) else "無刑罰"
    extraction_status = "完整" if raw_charges and any(raw_laws) and (declared or any(fines) or executions or execution_fines) else "需人工核對"
    row = [
        source.get("裁判字號", ""), source.get("裁判書連結", ""), source.get("被告／相對人", ""), defendant_status,
        source.get("裁判種類", ""), source.get("案件類型", ""), result, source.get("主文", ""),
        json_value(raw_charges), json_value(standard_charges), charge_source,
        json_value(raw_laws), json_value(standard_laws), json_value(law_sources),
        json_value(primaries), json_value(primary_months), json_value(fines),
        json_value(executions), json_value(sentence_months(executions)), json_value(execution_fines),
        json_value(selected), json_value(sentence_months(selected)),
        sentence_source, risk, sentence_status, extraction_status,
        raw_missing, standard_missing, month_missing, sentence_missing,
    ]
    denominator = result == "有罪" and defendant_status == "單被告"
    flags = {
        "denominator": denominator,
        "single_execution": denominator and sentence_source == "應執行刑",
        "single_declared": denominator and sentence_source == "宣告刑",
        "non_single": denominator and not sentence_source,
        "standard_gap": denominator and bool(standard_missing),
        "raw_gap": denominator and (charge_missing or law_missing),
        "month_gap": denominator and bool(month_missing),
        "sentence_gap": denominator and bool(sentence_missing),
    }
    return row, flags


# Build to a temporary path and replace the canonical workbook only after a complete save.
def build_workbook(source: Path, output: Path) -> dict[str, int]:
    """Extract all cases and atomically replace the canonical compact workbook."""
    source_book = load_workbook(source, read_only=True, data_only=True)
    if SOURCE_SHEET not in source_book.sheetnames:
        source_book.close()
        raise ValueError(f"找不到工作表：{SOURCE_SHEET}")
    source_sheet = source_book[SOURCE_SHEET]
    source_headers = [cell.value for cell in next(source_sheet.iter_rows(max_row=1))]
    output_book = Workbook(write_only=True)
    sheets = {name: add_sheet(output_book, name) for name in SHEETS}
    counts = {name: 0 for name in SHEETS}
    routing = {
        "denominator": "有罪_單被告", "single_execution": "單一應執行刑",
        "single_declared": "單一宣告刑", "non_single": "非單一刑期",
        "standard_gap": "罪法_有值但標準化無值", "raw_gap": "罪法_原始無值",
        "month_gap": "刑期_有值但月數無值", "sentence_gap": "刑期_原始無值",
    }
    for values in source_sheet.iter_rows(min_row=2, values_only=True):
        row, flags = build_row(dict(zip(source_headers, values)))
        sheets["罪法刑"].append(row)
        counts["罪法刑"] += 1
        for flag, sheet_name in routing.items():
            if flags[flag]:
                sheets[sheet_name].append(row)
                counts[sheet_name] += 1
    temporary = output.with_suffix(".updating.xlsx")
    output_book.save(temporary)
    source_book.close()
    os.replace(temporary, output)
    return counts


# Expose a repeatable command instead of keeping the final-data build as a one-off notebook cell.
def main() -> None:
    """Parse input and output paths, then print sheet counts after a successful build."""
    parser = argparse.ArgumentParser(description="建立罪法刑唯一審查檔")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if not args.input.exists():
        parser.error(f"找不到輸入檔：{args.input}")
    print(build_workbook(args.input, args.output))


if __name__ == "__main__":
    main()
